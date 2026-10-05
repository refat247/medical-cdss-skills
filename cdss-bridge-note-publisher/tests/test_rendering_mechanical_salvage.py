"""Rendering/image-only regression coverage salvaged from PR #1.

Grounding-policy, citation adjudication, and clinical claim-gating changes are
intentionally excluded from this mechanical rendering suite.
"""
import os
import sys
from pathlib import Path

import pytest

SCRIPTS = Path(__file__).resolve().parent.parent / "scripts"
sys.path.insert(0, str(SCRIPTS))

pytest.importorskip("docx")
import docx  # noqa: E402

from enhance_figures import audit_and_enhance_directory, enhance_figure  # noqa: E402
from publish_executive_docx import (  # noqa: E402
    _resolve_publication_image,
    add_formatted_runs,
    compile_executive_docx,
)


def build_docx(tmp_path, md_text):
    md = tmp_path / "n.md"
    md.write_text(md_text, encoding="utf-8")
    out = tmp_path / "n.docx"
    compile_executive_docx(md, out, base_dir=tmp_path, allow_missing_images=True)
    return docx.Document(str(out))


def test_docx_title_comes_from_note(tmp_path):
    document = build_docx(
        tmp_path,
        "# Diabetic Ketoacidosis\n\n## LAYER 2\n- Source-preserved note content.\n",
    )
    texts = [p.text for p in document.paragraphs]
    assert any("Diabetic Ketoacidosis" in t for t in texts)
    assert not any("Cardiac Auscultation" in t for t in texts)


def test_two_tables_separated_by_blank_line_stay_separate(tmp_path):
    md = "# T\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n| X | Y | Z |\n|---|---|---|\n| low | replace | keep |\n"
    document = build_docx(tmp_path, md)
    assert len(document.tables) == 2
    cells = [c.text for r in document.tables[1].rows for c in r.cells]
    assert "low" in cells and "replace" in cells and "keep" in cells


def test_escaped_pipe_does_not_create_extra_cell(tmp_path):
    document = build_docx(tmp_path, "# T\n\n| Formula | Note |\n|---|---|\n| P(A\\|B) | y |\n")
    row = [c.text for c in document.tables[0].rows[1].cells]
    assert row == ["P(A|B)", "y"]


def test_literal_asterisks_are_not_italic_markup():
    document = docx.Document()
    para = document.add_paragraph()
    add_formatted_runs(para, "dose a * b * c and *real italic* and **bold**")
    assert "".join(r.text for r in para.runs) == "dose a * b * c and real italic and bold"
    assert [r.text for r in para.runs if r.italic] == ["real italic"]


def test_br_tags_become_line_breaks():
    para = docx.Document().add_paragraph()
    add_formatted_runs(para, "first**bold**<br>second<br/>third")
    assert "<br" not in para.text
    assert para.text.replace("\n", "|") == "firstbold|second|third"


def test_separate_numbered_lists_restart_at_one(tmp_path):
    document = build_docx(tmp_path, "# T\n\n1. a\n2. b\n3. c\n\nBetween.\n\n1. x\n2. y\n")
    items = [p for p in document.paragraphs if p.style.name == "List Number"]
    ids = [p._p.pPr.numPr.numId.val for p in items]
    assert len(items) == 5
    assert len(set(ids[:3])) == 1
    assert len(set(ids[3:])) == 1
    assert ids[0] != ids[3]
    numbering = document.part.numbering_part.numbering_definitions._numbering
    for nid in (ids[0], ids[3]):
        override = numbering.num_having_numId(nid).lvlOverride_lst[0]
        assert override.startOverride.val == 1


def test_16bit_greyscale_png_is_not_whitened(tmp_path):
    PIL = pytest.importorskip("PIL.Image")
    src, dst = tmp_path / "g.png", tmp_path / "out.png"
    PIL.new("I;16", (8, 8), 30000).save(src)
    assert enhance_figure(src, dst, scale_factor=1)
    px = PIL.open(dst).convert("RGB").getpixel((4, 4))
    assert 100 < px[0] < 135


def test_enhanced_copy_refreshes_when_source_changes(tmp_path):
    PIL = pytest.importorskip("PIL.Image")
    figs, out = tmp_path / "figs", tmp_path / "out"
    figs.mkdir()
    PIL.new("RGB", (700, 10), (255, 0, 0)).save(figs / "a.png")
    audit_and_enhance_directory(figs, out)
    assert PIL.open(out / "a.png").getpixel((0, 0)) == (255, 0, 0)

    PIL.new("RGB", (700, 10), (0, 0, 255)).save(figs / "a.png")
    future = (out / "a.png").stat().st_mtime + 10
    os.utime(figs / "a.png", (future, future))
    audit_and_enhance_directory(figs, out)
    assert PIL.open(out / "a.png").getpixel((0, 0)) == (0, 0, 255)


def test_enhanced_lookup_does_not_alias_duplicate_basenames(tmp_path):
    a = tmp_path / "a" / "same.png"
    b = tmp_path / "b" / "same.png"
    a.parent.mkdir()
    b.parent.mkdir()
    a.write_bytes(b"A")
    b.write_bytes(b"B")
    enhanced = tmp_path / "figures_enhanced"
    enhanced.mkdir()
    (enhanced / "same.png").write_bytes(b"AMBIGUOUS")

    counts = {"same.png": 2}
    assert _resolve_publication_image(tmp_path, "a/same.png", counts) == a.resolve()
    assert _resolve_publication_image(tmp_path, "b/same.png", counts) == b.resolve()


def test_enhanced_lookup_keeps_unique_flat_backward_compatibility(tmp_path):
    src = tmp_path / "assets" / "only.png"
    src.parent.mkdir()
    src.write_bytes(b"SRC")
    enhanced = tmp_path / "figures_enhanced"
    enhanced.mkdir()
    flat = enhanced / "only.png"
    flat.write_bytes(b"ENHANCED")
    assert _resolve_publication_image(tmp_path, "assets/only.png", {"only.png": 1}) == flat.resolve()


def test_enhanced_lookup_prefers_mirrored_path(tmp_path):
    src = tmp_path / "a" / "same.png"
    src.parent.mkdir()
    src.write_bytes(b"SRC")
    mirrored = tmp_path / "figures_enhanced" / "a" / "same.png"
    mirrored.parent.mkdir(parents=True)
    mirrored.write_bytes(b"ENHANCED")
    assert _resolve_publication_image(tmp_path, "a/same.png", {"same.png": 2}) == mirrored.resolve()


def test_auscultation_code_block_preserves_note_text(tmp_path):
    md = (
        "# Custom Auscultation Note\n\n"
        "```\n"
        "7-ATTRIBUTE AUSCULTATION FRAMEWORK\n"
        "1. TIMING — custom evidence-locked timing text\n"
        "7. DYNAMIC — custom evidence-locked manoeuvre text\n"
        "```\n"
    )
    document = build_docx(tmp_path, md)
    text = "\n".join(p.text for p in document.paragraphs)
    assert "custom evidence-locked timing text" in text
    assert "custom evidence-locked manoeuvre text" in text
    assert "Carvallo sign" not in text
    assert "Brock-Braunwald" not in text
