"""Regression tests for the grounding-gate fail-open paths (skill_audits/SECOND_SWEEP.md 2.14-2.17)."""
import sys
from pathlib import Path

import pytest

SCRIPTS = Path(__file__).resolve().parent.parent / "scripts"
sys.path.insert(0, str(SCRIPTS))
from verify_grounding import verify_note_grounding  # noqa: E402


def make_package(root: Path) -> Path:
    pkg = root / "pkg"
    for ch, body in (("16", "Grade 4 has a thrill."), ("17", "Heart failure drugs.")):
        d = pkg / "01_Davidson_25" / f"Davidson_25_Ch{ch}_X"
        d.mkdir(parents=True)
        # local chunk ids repeat in every chapter, exactly like the real pipeline output
        (d / f"Davidson_25_Ch{ch}_X_RAG_Optimised.md").write_text(
            f"---\nchunk_id: L2-001\n---\n{body}\n---\nchunk_id: L2-002\n---\nMore text {ch}.\n", encoding="utf-8")
    # an id that exists in ONE chapter only
    (pkg / "01_Davidson_25" / "Davidson_25_Ch16_X" / "Davidson_25_Ch16_X_RAG_Optimised.md").write_text(
        "---\nchunk_id: L2-001\n---\nGrade 4 has a thrill.\n---\nchunk_id: L2-002\n---\nMore text 16.\n"
        "---\nchunk_id: L2-777\n---\nOnly in sixteen.\n", encoding="utf-8")
    return pkg


def run(tmp_path, text):
    note = tmp_path / "note.md"
    note.write_text(text, encoding="utf-8")
    return verify_note_grounding(note, make_package(tmp_path))


HEAD = "# T\n## LAYER 2 — Davidson Core Spine\n- Levine grade 4 has a thrill [Anchor: Davidson-16-L2-001]\n"


def test_baseline_passes(tmp_path):
    assert run(tmp_path, HEAD)["status"] == "PASS"


def test_exempt_words_in_a_layer_heading_do_not_exempt_the_section(tmp_path):
    rep = run(tmp_path, HEAD + "## Layer 3 - Pharmacology: Clinical Evidence\n- Digoxin 5 mg daily is first line for all atrial fibrillation patients.\n")
    assert rep["status"] == "FAIL" and rep["uncited_lines"]


def test_real_exempt_section_is_still_exempt(tmp_path):
    rep = run(tmp_path, HEAD + "## Clinical Evidence Notes\n- ESC 2021 gives a Class I recommendation for echocardiography here.\n")
    assert rep["status"] == "PASS"


def test_unclassified_section_with_claims_fails(tmp_path):
    rep = run(tmp_path, HEAD + "## Treatment Pearls\n- Amiodarone 2000 mg IV bolus is safe in all patients with shock.\n")
    assert rep["status"] == "FAIL"


def test_bare_chunk_id_that_exists_in_several_chapters_is_ambiguous(tmp_path):
    rep = run(tmp_path, "# T\n## LAYER 2 — Davidson Core Spine\n- Amiodarone 2000 mg IV bolus in all patients [chunk: L2-001]\n")
    assert rep["status"] == "FAIL" and any("ambiguous" in u.lower() for u in rep["unresolved_citations"])


def test_unique_bare_chunk_id_still_resolves(tmp_path):
    rep = run(tmp_path, "# T\n## LAYER 2 — Davidson Core Spine\n- Only in sixteen exactly as written here [chunk: L2-777]\n")
    assert rep["status"] == "PASS", rep["failures"]


def test_wrong_chapter_label_fails_even_for_a_unique_id(tmp_path):
    rep = run(tmp_path, "# T\n## LAYER 2 — Davidson Core Spine\n- Only in sixteen but cited as ninety nine [Anchor: Davidson-99-L2-777]\n")
    assert rep["status"] == "FAIL"


def test_single_token_table_rows_are_claims(tmp_path):
    rep = run(tmp_path, HEAD + "## LAYER 4 — Beyond Davidson\n| Drug | Dose |\n|---|---|\n| Warfarin | 10mg |\n| Digoxin | 5mg |\n")
    assert rep["status"] == "FAIL" and len(rep["uncited_lines"]) == 2


def test_layer_1_heading_does_not_match_layer_10(tmp_path):
    rep = run(tmp_path, HEAD + "## Layer 10 notes\n- This line has no citation but is plainly a clinical claim here.\n")
    assert rep["status"] == "FAIL"        # unclassified (not silently treated as Layer 1 or exempt)


# ---------- docx publisher: title, tables ----------
docx = pytest.importorskip("docx")
from publish_executive_docx import compile_executive_docx  # noqa: E402


def build_docx(tmp_path, md_text):
    md = tmp_path / "n.md"; md.write_text(md_text, encoding="utf-8")
    out = tmp_path / "n.docx"
    compile_executive_docx(md, out, base_dir=tmp_path, allow_missing_images=True)
    return docx.Document(str(out))


def test_docx_title_comes_from_the_note_not_a_hard_coded_cardiac_title(tmp_path):
    d = build_docx(tmp_path, "# Diabetic Ketoacidosis\n\n## LAYER 2\n- Insulin infusion is started [Anchor: Davidson-16-L2-001]\n")
    texts = [p.text for p in d.paragraphs]
    assert any("Diabetic Ketoacidosis" in t for t in texts)
    assert not any("Cardiac Auscultation" in t for t in texts)


def test_two_tables_separated_by_a_blank_line_stay_two_tables(tmp_path):
    md = "# T\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n| X | Y | Z |\n|---|---|---|\n| low | replace | keep |\n"
    d = build_docx(tmp_path, md)
    assert len(d.tables) == 2
    cells = [c.text for r in d.tables[1].rows for c in r.cells]
    assert "low" in cells and "replace" in cells and "keep" in cells


def test_escaped_pipe_does_not_create_an_extra_cell(tmp_path):
    md = "# T\n\n| Formula | Note |\n|---|---|\n| P(A\\|B) | y |\n"
    d = build_docx(tmp_path, md)
    row = [c.text for c in d.tables[0].rows[1].cells]
    assert row == ["P(A|B)", "y"]


def test_literal_asterisks_are_not_treated_as_italic(tmp_path):
    """M28: 'a * b * c' lost its asterisks and italicised ' b '."""
    pytest.importorskip("docx")
    import docx
    from publish_executive_docx import add_formatted_runs
    doc = docx.Document()
    para = doc.add_paragraph()
    add_formatted_runs(para, "dose a * b * c and *real italic* and **bold**")
    assert "".join(r.text for r in para.runs) == "dose a * b * c and real italic and bold"
    italic = [r.text for r in para.runs if r.italic]
    assert italic == ["real italic"]


def test_16bit_greyscale_png_is_not_whitened(tmp_path):
    """M30: I;16 mid-grey converted to pure white."""
    PIL = pytest.importorskip("PIL.Image")
    from enhance_figures import enhance_figure
    src, dst = tmp_path / "g.png", tmp_path / "out.png"
    PIL.new("I;16", (8, 8), 30000).save(src)
    assert enhance_figure(src, dst, scale_factor=1)
    px = PIL.open(dst).convert("RGB").getpixel((4, 4))
    assert 100 < px[0] < 135                       # ~30000/256 = 117, not 255


def test_enhanced_copy_is_refreshed_when_source_changes(tmp_path):
    """M31: an existing output copy was never replaced after the source figure changed."""
    import os
    PIL = pytest.importorskip("PIL.Image")
    from enhance_figures import audit_and_enhance_directory
    figs, out = tmp_path / "figs", tmp_path / "out"
    figs.mkdir()
    PIL.new("RGB", (700, 10), (255, 0, 0)).save(figs / "a.png")
    audit_and_enhance_directory(figs, out)
    assert PIL.open(out / "a.png").getpixel((0, 0)) == (255, 0, 0)
    PIL.new("RGB", (700, 10), (0, 0, 255)).save(figs / "a.png")
    future = (out / "a.png").stat().st_mtime + 10
    os.utime(figs / "a.png", (future, future))
    audit_and_enhance_directory(figs, out)
    assert PIL.open(out / "a.png").getpixel((0, 0)) == (0, 0, 255)


def test_book_and_chapter_matching_are_not_loose():
    """M32: 'Research_12' parsed as chapter 12; a name mentioning two books resolved to the first."""
    from verify_grounding import CHAPTER_RE, normalise_book
    assert CHAPTER_RE.search("Research_12") is None and CHAPTER_RE.search("Search 3") is None
    assert CHAPTER_RE.search("Davidson_25_Ch16_X").group(1) == "16"
    assert CHAPTER_RE.search("Chapter 7").group(1) == "7"
    assert normalise_book("Davidson_vs_Harrison_notes") is None
    assert normalise_book("01_Davidson_25") == "davidson"


def test_br_tags_become_line_breaks_not_literal_text():
    """M28: '<br>' was printed literally in table cells."""
    pytest.importorskip("docx")
    import docx
    from publish_executive_docx import add_formatted_runs
    para = docx.Document().add_paragraph()
    add_formatted_runs(para, "first**bold**<br>second<br/>third")
    assert "<br" not in para.text
    assert para.text.replace("\n", "|") == "firstbold|second|third"


def test_m29_separate_numbered_lists_each_restart_at_one(tmp_path):
    """M29: every 'List Number' paragraph shared one numId, so a second list carried on from the first."""
    pytest.importorskip("docx")
    import docx
    from publish_executive_docx import compile_executive_docx
    md = tmp_path / "n.md"
    md.write_text("# T\n\n1. a\n2. b\n3. c\n\nBetween the lists.\n\n1. x\n2. y\n", encoding="utf-8")
    out = tmp_path / "n.docx"
    compile_executive_docx(md, out)
    d = docx.Document(str(out))
    items = [p for p in d.paragraphs if p.style.name == "List Number"]
    ids = [p._p.pPr.numPr.numId.val if p._p.pPr is not None and p._p.pPr.numPr is not None else None for p in items]
    assert len(items) == 5 and None not in ids
    assert len(set(ids[:3])) == 1 and len(set(ids[3:])) == 1 and ids[0] != ids[3]   # one num per list
    numbering = d.part.numbering_part.numbering_definitions._numbering
    for nid in (ids[0], ids[3]):
        override = numbering.num_having_numId(nid).lvlOverride_lst[0]
        assert override.startOverride.val == 1



def test_bridge_enhanced_lookup_does_not_alias_duplicate_basenames(tmp_path):
    from publish_executive_docx import _resolve_publication_image

    a = tmp_path / "a" / "same.png"
    b = tmp_path / "b" / "same.png"
    a.parent.mkdir()
    b.parent.mkdir()
    a.write_bytes(b"A")
    b.write_bytes(b"B")

    enhanced = tmp_path / "figures_enhanced"
    enhanced.mkdir()
    (enhanced / "same.png").write_bytes(b"AMBIGUOUS-FLAT")

    counts = {"same.png": 2}
    assert _resolve_publication_image(tmp_path, "a/same.png", counts) == a.resolve()
    assert _resolve_publication_image(tmp_path, "b/same.png", counts) == b.resolve()


def test_bridge_enhanced_lookup_keeps_unique_flat_backward_compatibility(tmp_path):
    from publish_executive_docx import _resolve_publication_image

    src = tmp_path / "assets" / "only.png"
    src.parent.mkdir()
    src.write_bytes(b"SRC")
    enhanced = tmp_path / "figures_enhanced"
    enhanced.mkdir()
    flat = enhanced / "only.png"
    flat.write_bytes(b"ENHANCED")

    assert _resolve_publication_image(tmp_path, "assets/only.png", {"only.png": 1}) == flat.resolve()


def test_bridge_enhanced_lookup_prefers_mirrored_path_when_available(tmp_path):
    from publish_executive_docx import _resolve_publication_image

    src = tmp_path / "a" / "same.png"
    src.parent.mkdir()
    src.write_bytes(b"SRC")
    mirrored = tmp_path / "figures_enhanced" / "a" / "same.png"
    mirrored.parent.mkdir(parents=True)
    mirrored.write_bytes(b"ENHANCED")

    assert _resolve_publication_image(tmp_path, "a/same.png", {"same.png": 2}) == mirrored.resolve()



def test_auscultation_framework_code_block_preserves_note_text_not_canned_content(tmp_path):
    md = (
        "# Custom Auscultation Note\n\n"
        "```\n"
        "7-ATTRIBUTE AUSCULTATION FRAMEWORK\n"
        "1. TIMING — custom evidence-locked timing text\n"
        "7. DYNAMIC — custom evidence-locked manoeuvre text\n"
        "```\n"
    )
    d = build_docx(tmp_path, md)
    text = "\n".join(p.text for p in d.paragraphs)
    assert "custom evidence-locked timing text" in text
    assert "custom evidence-locked manoeuvre text" in text
    assert "Carvallo sign" not in text
    assert "Brock-Braunwald" not in text
