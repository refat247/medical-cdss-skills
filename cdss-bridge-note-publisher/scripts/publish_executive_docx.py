"""
CDSS Executive Word (.docx) Publisher.
Compiles clinical bridge notes into journal-grade Word documents featuring:
- Native Word Card Grid Tables (zero ASCII staircase art)
- Border-accented Alert Cards (Coverage, QB Trap, Quick Revision)
- 4x Lanczos-enhanced figures with adaptive column widths
- Clinical tables with executive styling
"""

import re
import sys
import argparse
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from PIL import Image

# Palette
COLOR_PRIMARY = RGBColor(10, 37, 64)       # Deep Navy
COLOR_SECONDARY = RGBColor(25, 103, 210)   # Sapphire Blue
COLOR_ACCENT = RGBColor(0, 150, 136)       # Medical Teal
COLOR_TEXT_DARK = RGBColor(32, 33, 36)     # Off-Black Body
COLOR_MUTED = RGBColor(95, 99, 104)        # Slate Grey
COLOR_QB = RGBColor(197, 34, 31)           # Crimson Trap
COLOR_WHITE = RGBColor(255, 255, 255)

HEX_PRIMARY = "0A2540"
HEX_ROW_ALT = "F8FAFC"
HEX_ROW_WHITE = "FFFFFF"
HEX_BORDER = "CBD5E1"
HEX_CARD_BG = "F8FAFC"
HEX_CARD_ALT = "F1F5F9"
HEX_QB_BG = "FFF8E1"
HEX_COV_BG = "F0F4F8"
HEX_REV_BG = "F6FBF7"

PRINTABLE_WIDTH_INCHES = 6.5


def set_cell_background(cell, hex_color: str):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_table_borders(table, color="CBD5E1", sz="4"):
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="0A2540"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="0A2540"/>'
        f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    table._tbl.tblPr.append(tblBorders)


def make_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))


def restart_numbering(doc, paragraph):
    """Give this list and its following contiguous items a fresh numbering instance starting at 1."""
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    style_num_id = paragraph.style.element.pPr.numPr.numId.val
    abstract_id = numbering.num_having_numId(style_num_id).abstractNumId.val
    num = numbering.add_num(abstract_id)
    num.add_lvlOverride(ilvl=0).add_startOverride(1)
    paragraph._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = num.numId


def add_formatted_runs(paragraph, text: str, default_font="Calibri", default_size=11, default_color=COLOR_TEXT_DARK):
    if not text:
        return
    br_parts = re.split(r'<br\s*/?>', text, flags=re.IGNORECASE)
    if len(br_parts) > 1:
        for n, part in enumerate(br_parts):
            if n:
                paragraph.add_run().add_break()
            add_formatted_runs(paragraph, part.strip(), default_font, default_size, default_color)
        return
    # Emphasis markers must hug non-space content; literal arithmetic-like asterisks remain literal.
    pattern = re.compile(
        r'(\*\*\*(?=\S).*?(?<=\S)\*\*\*|\*\*(?=\S).*?(?<=\S)\*\*|'
        r'\*(?=[^\s*]).*?(?<=[^\s*])\*|`.*?`|\[.*?\]\(.*?\))'
    )
    tokens = pattern.split(text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        run.font.name = default_font
        run.font.size = Pt(default_size)
        run.font.color.rgb = default_color
        if token.startswith('***') and token.endswith('***') and len(token) >= 6:
            run.text = token[3:-3]
            run.bold = True
            run.italic = True
        elif token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run.text = token[2:-2]
            run.bold = True
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run.text = token[1:-1]
            run.italic = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run.text = token[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(default_size - 1)
            run.font.color.rgb = COLOR_SECONDARY
        elif token.startswith('[') and '](' in token and token.endswith(')'):
            m = re.match(r'\[(.*?)\]\((.*?)\)', token)
            if m:
                run.text = m.group(1)
                run.font.color.rgb = COLOR_SECONDARY
                run.underline = True
            else:
                run.text = token
        else:
            run.text = token


def create_card_grid_table(doc, title: str, row1_cards: list, row2_cards: list):
    """Builds a journal-grade 4-column card grid table in Word."""
    tbl = doc.add_table(rows=3, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = tbl.rows[0].cells[0]
    for c in tbl.rows[0].cells[1:]:
        hdr.merge(c)

    set_cell_background(hdr, HEX_PRIMARY)
    set_cell_margins(hdr, top=140, bottom=140, left=160, right=160)
    p = hdr.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = COLOR_WHITE

    # Row 1 Cards (4 Columns)
    for idx, (c_title, items) in enumerate(row1_cards[:4]):
        cell = tbl.rows[1].cells[idx]
        cell.width = Inches(1.625)
        set_cell_background(cell, HEX_CARD_BG)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r_title = p.add_run(c_title)
        r_title.font.name = "Arial"
        r_title.font.size = Pt(9.5)
        r_title.font.bold = True
        r_title.font.color.rgb = COLOR_SECONDARY
        for item in items:
            pi = cell.add_paragraph()
            pi.paragraph_format.space_before = Pt(1)
            pi.paragraph_format.space_after = Pt(1)
            pi.paragraph_format.line_spacing = 1.05
            add_formatted_runs(pi, f"• {item}", default_size=8.5)

    # Row 2 Cards
    row2_cells = tbl.rows[2].cells
    if len(row2_cards) == 3:
        row2_cells[2].merge(row2_cells[3])
        specs = [
            (row2_cells[0], Inches(1.625), row2_cards[0][0], row2_cards[0][1]),
            (row2_cells[1], Inches(1.625), row2_cards[1][0], row2_cards[1][1]),
            (row2_cells[2], Inches(3.25), row2_cards[2][0], row2_cards[2][1]),
        ]
    else:
        specs = [(row2_cells[i], Inches(1.625), row2_cards[i][0], row2_cards[i][1]) for i in range(len(row2_cards))]

    for cell, width, c_title, items in specs:
        cell.width = width
        set_cell_background(cell, HEX_CARD_ALT)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r_title = p.add_run(c_title)
        r_title.font.name = "Arial"
        r_title.font.size = Pt(9.5)
        r_title.font.bold = True
        r_title.font.color.rgb = COLOR_PRIMARY
        for item in items:
            pi = cell.add_paragraph()
            pi.paragraph_format.space_before = Pt(1)
            pi.paragraph_format.space_after = Pt(1)
            pi.paragraph_format.line_spacing = 1.05
            add_formatted_runs(pi, f"• {item}", default_size=8.5)

    set_table_borders(tbl)
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(2)
    p_after.paragraph_format.space_after = Pt(8)


def create_callout_box(doc, text_content: str, box_type: str = "cov"):
    """Renders a single-cell bordered callout card with colored accent bar."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(PRINTABLE_WIDTH_INCHES)

    if box_type == "qb":
        bg_hex = HEX_QB_BG
        border_hex = "D93025"
        title_text = "CRITICAL EXAM TRAP / QUESTION BANK AWARENESS (DUAL-TRUTH EXCEPTION LAYER)"
        title_color = COLOR_QB
    elif box_type == "rev":
        bg_hex = HEX_REV_BG
        border_hex = "137333"
        title_text = "HIGH-YIELD RAPID REVISION CHECKLIST"
        title_color = RGBColor(19, 115, 51)
    else:
        bg_hex = HEX_COV_BG
        border_hex = "0A2540"
        title_text = "V2.2 COVERAGE DECLARATION & DUAL-TRUTH ARCHITECTURE"
        title_color = COLOR_PRIMARY

    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)

    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
        f'  <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'</w:tblBorders>'
    )
    tbl._tbl.tblPr.append(tblBorders)

    tp = cell.paragraphs[0]
    tp.paragraph_format.space_before = Pt(2)
    tp.paragraph_format.space_after = Pt(4)
    tr = tp.add_run(title_text)
    tr.font.name = "Arial"
    tr.font.size = Pt(10)
    tr.font.bold = True
    tr.font.color.rgb = title_color

    clean_lines = []
    for line in text_content.splitlines():
        line = line.strip()
        if not line or line.startswith("+---") or line.startswith("|==="):
            continue
        line = line.strip("|").strip()
        if line.startswith("[!NOTE]") or line.startswith("[!WARNING]") or line.startswith("[!TIP]"):
            continue
        if "COVERAGE DECLARATION" in line or "QB AWARENESS" in line or "QUICK REVISION" in line:
            continue
        clean_lines.append(line)

    for cline in clean_lines:
        cp = cell.add_paragraph()
        cp.paragraph_format.space_before = Pt(1)
        cp.paragraph_format.space_after = Pt(2)
        cp.paragraph_format.line_spacing = 1.05
        if cline.startswith("* ") or cline.startswith("- "):
            cp.paragraph_format.left_indent = Inches(0.12)
        elif re.match(r'^\d+\.\s+', cline):
            cp.paragraph_format.left_indent = Inches(0.12)
        add_formatted_runs(cp, cline, default_size=9)

    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_before = Pt(2)
    p_post.paragraph_format.space_after = Pt(6)


def _split_table_row(line: str):
    """Split a markdown table row while preserving escaped pipes (\|) inside cells."""
    s = line.strip()
    cells = re.split(r"(?<!\\)\|", s)
    if s.startswith("|"):
        cells = cells[1:]
    if s.endswith("|") and not s.endswith("\\|"):
        cells = cells[:-1]
    return [c.replace("\\|", "|").strip() for c in cells]


def _resolve_publication_image(base_dir: Path, rel_path: str, basename_counts: dict) -> Path:
    """Resolve enhanced figures without aliasing distinct source images that share a basename."""
    img_file = (base_dir / rel_path).resolve()
    enhanced_root = (base_dir / "figures_enhanced").resolve()
    rel = Path(rel_path)
    mirrored = (enhanced_root / rel).resolve()

    try:
        mirrored.relative_to(enhanced_root)
    except ValueError:
        mirrored = enhanced_root / "__invalid__"

    if mirrored.exists():
        return mirrored

    flat = enhanced_root / img_file.name
    if basename_counts.get(img_file.name, 0) == 1 and flat.exists():
        return flat
    return img_file


def compile_executive_docx(md_path: Path, docx_path: Path, base_dir=None, allow_missing_images=False):
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    hdr = doc.sections[0].header
    hp = hdr.paragraphs[0]
    hp.text = "Clinical Knowledge System  |  Davidson Bridge Note V2.2 (Executive Edition)"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.runs[0].font.name = "Calibri"
    hp.runs[0].font.size = Pt(8.5)
    hp.runs[0].font.color.rgb = COLOR_MUTED

    tp = doc.add_paragraph()
    tp.paragraph_format.space_before = Pt(12)
    tp.paragraph_format.space_after = Pt(2)
    h1 = next((re.match(r"^#\s+(.*\S)\s*$", l).group(1) for l in lines if re.match(r"^#\s+\S", l)), None)
    note_title = re.sub(r"[*_`]+", "", h1) if h1 else md_path.stem.replace("_", " ")
    tr = tp.add_run(note_title)
    tr.font.name = "Arial"
    tr.font.size = Pt(24)
    tr.font.bold = True
    tr.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    srun = sub_p.add_run("Executive Davidson Cognitive Bridge Note V2.2  •  Native Vector Architecture")
    srun.font.name = "Calibri"
    srun.font.size = Pt(12)
    srun.font.italic = True
    srun.font.color.rgb = COLOR_SECONDARY

    base_dir = Path(base_dir) if base_dir else md_path.parent
    missing_images = []
    image_basename_counts = {}
    for source_line in lines:
        m = re.match(r'^!\[(.*?)\]\((.*?)\)', source_line.strip())
        if m:
            name = Path(m.group(2).strip()).name
            image_basename_counts[name] = image_basename_counts.get(name, 0) + 1

    i = 0
    total_lines = len(lines)

    while i < total_lines:
        line = lines[i]
        stripped = line.strip()

        if not stripped or stripped == "---":
            i += 1
            continue

        if stripped.startswith(">"):
            bq_lines = []
            while i < total_lines and lines[i].strip().startswith(">"):
                raw_l = lines[i].strip()
                content_l = re.sub(r'^>\s?', '', raw_l)
                bq_lines.append(content_l)
                i += 1

            bq_text = "\n".join(bq_lines)
            if "[!WARNING]" in bq_text or "QB AWARENESS" in bq_text or "CRITICAL EXAM TRAP" in bq_text:
                create_callout_box(doc, bq_text, box_type="qb")
            elif "[!TIP]" in bq_text or "QUICK REVISION" in bq_text or "HIGH-YIELD" in bq_text:
                create_callout_box(doc, bq_text, box_type="rev")
            elif "[!NOTE]" in bq_text or "COVERAGE DECLARATION" in bq_text:
                create_callout_box(doc, bq_text, box_type="cov")
            else:
                for bq_l in bq_lines:
                    if not bq_l.strip():
                        continue
                    bp = doc.add_paragraph()
                    bp.paragraph_format.left_indent = Inches(0.25)
                    bp.paragraph_format.space_before = Pt(2)
                    bp.paragraph_format.space_after = Pt(4)
                    bp.paragraph_format.line_spacing = 1.1
                    add_formatted_runs(bp, bq_l.strip(), default_font="Calibri", default_size=10, default_color=COLOR_TEXT_DARK)
            continue

        if stripped.startswith("```"):
            i += 1
            code_lines = []
            while i < total_lines and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if i < total_lines and lines[i].strip().startswith("```"):
                i += 1

            code_text = "\n".join(code_lines)

            # Never replace source note text with a canned clinical framework.
            if "COVERAGE DECLARATION" in code_text:
                create_callout_box(doc, code_text, box_type="cov")
                continue
            elif "QB AWARENESS" in code_text:
                create_callout_box(doc, code_text, box_type="qb")
                continue
            elif "QUICK REVISION BOX" in code_text:
                create_callout_box(doc, code_text, box_type="rev")
                continue
            else:
                cp = doc.add_paragraph()
                cp.paragraph_format.left_indent = Inches(0.15)
                cp.paragraph_format.space_before = Pt(4)
                cp.paragraph_format.space_after = Pt(6)
                r = cp.add_run(code_text)
                r.font.name = "Consolas"
                r.font.size = Pt(8.5)
                continue

        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            alt_text = img_match.group(1).strip()
            rel_path = img_match.group(2).strip()
            target_img = _resolve_publication_image(base_dir, rel_path, image_basename_counts)

            if not target_img.exists():
                missing_images.append(rel_path)
            if target_img.exists():
                with Image.open(target_img) as pil_img:
                    w_px, _ = pil_img.size
                doc_w = Inches(5.4) if w_px >= 1000 else Inches(4.5) if w_px >= 700 else Inches(3.4) if w_px >= 500 else Inches(2.8)

                ip = doc.add_paragraph()
                ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
                ip.paragraph_format.space_before = Pt(12)
                ip.paragraph_format.space_after = Pt(3)
                ip.paragraph_format.keep_with_next = True
                irun = ip.add_run()
                irun.add_picture(str(target_img), width=doc_w)

                caption_text = alt_text
                k = i + 1
                while k < total_lines and not lines[k].strip():
                    k += 1
                if k < total_lines and lines[k].strip().startswith("*Figure"):
                    caption_text = lines[k].strip().strip("*")
                    i = k

                cap_p = doc.add_paragraph()
                cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap_p.paragraph_format.space_before = Pt(2)
                cap_p.paragraph_format.space_after = Pt(14)
                c_run = cap_p.add_run(caption_text)
                c_run.font.name = "Calibri"
                c_run.font.size = Pt(9.5)
                c_run.font.italic = True
                c_run.font.color.rgb = COLOR_MUTED
            i += 1
            continue

        # 7-Attribute Card Grid Table Interceptor (Markdown dual tables)
        if stripped.startswith("|") and "1. TIMING" in stripped:
            t1_lines = []
            while i < total_lines and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                t1_lines.append(lines[i].strip())
                i += 1
            while i < total_lines and not lines[i].strip():
                i += 1
            t2_lines = []
            while i < total_lines and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                t2_lines.append(lines[i].strip())
                i += 1

            def parse_card_table(t_lines):
                if len(t_lines) < 3:
                    return []
                headers = _split_table_row(t_lines[0])
                content_row = _split_table_row(t_lines[2])
                cards = []
                for h, c in zip(headers, content_row):
                    raw_items = [item.strip() for item in re.split(r'<br\s*/?>', c) if item.strip()]
                    clean_items = [it.lstrip("•").strip() for it in raw_items if it.strip()]
                    cards.append((h, clean_items))
                return cards

            r1 = parse_card_table(t1_lines)
            r2 = parse_card_table(t2_lines)
            if r1 and r2:
                create_card_grid_table(doc, "THE SYSTEMATIC 7-ATTRIBUTE AUSCULTATION FRAMEWORK", r1, r2)
            else:
                # Malformed special tables fall back to source-preserving text; do not inject canned clinical content.
                for raw in t1_lines + t2_lines:
                    p = doc.add_paragraph()
                    add_formatted_runs(p, raw, default_font="Consolas", default_size=8.5)
            continue

        # Standard Markdown Table
        if stripped.startswith("|") and stripped.endswith("|") and not stripped.startswith("|   7-"):
            table_lines = []
            j = i
            while j < total_lines:
                curr = lines[j].strip()
                if curr.startswith("|") and curr.endswith("|"):
                    table_lines.append(curr)
                    j += 1
                elif not curr:
                    k = j + 1
                    while k < total_lines and not lines[k].strip():
                        k += 1
                    starts_table = k < total_lines and lines[k].strip().startswith("|") and lines[k].strip().endswith("|")
                    new_table = starts_table and k + 1 < total_lines and re.match(r'^\|[\s\:\-\|]+\|$', lines[k + 1].strip())
                    if starts_table and not new_table:
                        j = k
                    else:
                        break
                else:
                    break

            if len(table_lines) >= 2:
                raw_header = _split_table_row(table_lines[0])
                num_cols = len(raw_header)
                data_rows = []
                for row_line in table_lines[1:]:
                    if re.match(r'^\|[\s\:\-\|]+\|$', row_line):
                        continue
                    cols = _split_table_row(row_line)
                    if len(cols) < num_cols:
                        cols += [""] * (num_cols - len(cols))
                    data_rows.append(cols[:num_cols])

                t_font = 8.0 if num_cols == 8 else 8.5 if num_cols == 5 else 9.0 if num_cols == 4 else 9.5
                col_widths = [Inches(1.4)] + [Inches(0.728)] * 7 if num_cols == 8 else \
                             [Inches(1.2), Inches(1.8), Inches(0.9), Inches(1.1), Inches(1.5)] if num_cols == 5 else \
                             [Inches(1.2), Inches(2.2), Inches(2.1), Inches(1.0)] if num_cols == 4 else \
                             [Inches(PRINTABLE_WIDTH_INCHES / num_cols)] * num_cols

                tbl = doc.add_table(rows=len(data_rows) + 1, cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(tbl)

                hdr_row = tbl.rows[0]
                make_row_header(hdr_row)
                for idx, text_val in enumerate(raw_header):
                    c = hdr_row.cells[idx]
                    c.width = col_widths[idx]
                    set_cell_background(c, HEX_PRIMARY)
                    set_cell_margins(c, top=120, bottom=120, left=140, right=140)
                    p = c.paragraphs[0]
                    p.paragraph_format.line_spacing = 1.0
                    r = p.add_run(text_val)
                    r.font.name = "Calibri"
                    r.font.size = Pt(t_font + 0.5)
                    r.font.bold = True
                    r.font.color.rgb = COLOR_WHITE

                for r_idx, r_data in enumerate(data_rows):
                    row = tbl.rows[r_idx + 1]
                    prevent_row_split(row)
                    bg = HEX_ROW_ALT if r_idx % 2 == 1 else HEX_ROW_WHITE
                    for c_idx, val in enumerate(r_data):
                        c = row.cells[c_idx]
                        c.width = col_widths[c_idx]
                        set_cell_background(c, bg)
                        set_cell_margins(c, top=80, bottom=80, left=140, right=140)
                        p = c.paragraphs[0]
                        p.paragraph_format.line_spacing = 1.05
                        add_formatted_runs(p, val, default_size=t_font)

                p_after = doc.add_paragraph()
                p_after.paragraph_format.space_after = Pt(6)
                i = j
                continue

        # Headings
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            htext = stripped.lstrip("#").strip()
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            run = p.add_run(htext)
            run.font.name = "Arial"
            run.font.bold = True
            if level == 1:
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(8)
                run.font.size = Pt(18)
                run.font.color.rgb = COLOR_PRIMARY
            elif level == 2:
                p.paragraph_format.space_before = Pt(16)
                p.paragraph_format.space_after = Pt(4)
                run.font.size = Pt(13.5)
                run.font.color.rgb = COLOR_PRIMARY
            elif level == 3:
                p.paragraph_format.space_before = Pt(11)
                p.paragraph_format.space_after = Pt(3)
                run.font.size = Pt(11.5)
                run.font.color.rgb = COLOR_SECONDARY
            else:
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(2)
                run.font.size = Pt(10.5)
                run.font.color.rgb = COLOR_TEXT_DARK
            i += 1
            continue

        # Lists & Paragraphs
        bullet_m = re.match(r'^(\s*)([\*\-])\s+(.*)$', line)
        if bullet_m:
            bp = doc.add_paragraph(style='List Bullet')
            bp.paragraph_format.space_before = Pt(1)
            bp.paragraph_format.space_after = Pt(2)
            bp.paragraph_format.line_spacing = 1.1
            bp.paragraph_format.left_indent = Inches(0.45 if len(bullet_m.group(1)) >= 2 else 0.25)
            add_formatted_runs(bp, bullet_m.group(3))
            i += 1
            continue

        num_m = re.match(r'^(\s*)(\d+)\.\s+(.*)$', line)
        if num_m:
            prev = doc.paragraphs[-1] if doc.paragraphs else None
            np_p = doc.add_paragraph(style='List Number')
            if prev is None or prev.style.name != 'List Number':
                restart_numbering(doc, np_p)
            else:
                prev_pPr = prev._p.pPr
                if prev_pPr is not None and prev_pPr.numPr is not None:
                    np_p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = prev_pPr.numPr.numId.val
            np_p.paragraph_format.space_before = Pt(1)
            np_p.paragraph_format.space_after = Pt(2)
            np_p.paragraph_format.line_spacing = 1.1
            np_p.paragraph_format.left_indent = Inches(0.45 if len(num_m.group(1)) >= 2 else 0.25)
            add_formatted_runs(np_p, num_m.group(3))
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_runs(p, stripped)
        i += 1

    if missing_images and not allow_missing_images:
        raise FileNotFoundError(f"{len(missing_images)} image(s) referenced in {md_path.name} not found under {base_dir}: {missing_images[:10]}")
    for m in missing_images:
        print(f"[PUBLISH-DOCX] WARNING: missing image skipped: {m}")
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    print(f"[PUBLISH-DOCX] Successfully saved: {docx_path}")


def main():
    parser = argparse.ArgumentParser(description="Compile Executive Word (.docx) Document")
    parser.add_argument("--input-md", "-i", required=True, help="Input Markdown file")
    parser.add_argument("--output-docx", "-o", required=True, help="Output Word .docx file")
    parser.add_argument("--base-dir", help="Folder that relative image paths resolve against (default: the markdown's folder)")
    parser.add_argument("--allow-missing-images", action="store_true", help="Publish even if referenced images are missing")
    args = parser.parse_args()
    try:
        compile_executive_docx(Path(args.input_md), Path(args.output_docx), base_dir=args.base_dir,
                               allow_missing_images=args.allow_missing_images)
    except FileNotFoundError as e:
        print(f"[PUBLISH-DOCX] ERROR: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
