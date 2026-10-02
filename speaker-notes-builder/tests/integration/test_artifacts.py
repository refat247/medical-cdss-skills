from pathlib import Path
import json

from pptx import Presentation
from pptx.util import Inches
from docx import Document

from speaker_notes_builder.docx_builder import build_docx
from speaker_notes_builder.qa import inject_notes, verify_pptx
from speaker_notes_builder.core import sha256_file


def _make_pptx(path: Path):
    prs = Presentation()
    layout = prs.slide_layouts[6]
    for i in range(2):
        slide = prs.slides.add_slide(layout)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
        box.text = f"Slide {i+1} title\nClinical content {i+1}"
    # remove initial default slide if any layout created none; Presentation() starts with zero slides
    prs.save(path)


def test_docx_build_and_note_injection(tmp_path):
    project = tmp_path / "project"
    project.mkdir()
    (project / "source").mkdir(); (project / "slides").mkdir(); (project / "final").mkdir()
    pptx = project / "source" / "source.pptx"
    _make_pptx(pptx)

    # lightweight slide images
    from PIL import Image, ImageDraw
    slides = []
    for i in (1, 2):
        img = project / "slides" / f"slide-{i:03d}.png"
        im = Image.new("RGB", (800, 450), "white")
        ImageDraw.Draw(im).text((50,50), f"Slide {i}", fill="black")
        im.save(img)
        slides.append({"slide": i, "title": f"Slide {i}", "image": str(img.relative_to(project)), "include": True, "role_guess": "general"})

    manifest = {"source_type": "pptx", "source_file": "source/source.pptx", "source_sha256": sha256_file(pptx), "slide_count": 2, "slides": slides, "duplicates": []}
    (project / "deck_manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    scripts = {"slides": []}
    for i in (1, 2):
        scripts["slides"].append({
            "slide": i, "title": f"Slide {i}", "section": "Test", "included": True,
            "delivery_intent": "Explain the test slide.",
            "main_script": f"This is the main script for slide {i}.",
            "compressed_script": f"Short script {i}.",
            "optional_expansion": f"Optional expansion {i}.",
            "key_emphasis": ["Key point"], "pause_or_point": ["Pause"], "transition": "Continue.",
            "estimated_seconds": 20, "cautions": [], "claim_status": "SUPPORTED"
        })
    (project / "scripts.json").write_text(json.dumps(scripts), encoding="utf-8")

    rehearsal = build_docx(project, "rehearsal")
    live = build_docx(project, "live")
    assert rehearsal.exists() and live.exists()
    assert len(Document(rehearsal).paragraphs) > 0

    out = project / "final" / "with_notes.pptx"
    inject_notes(project, out)
    result = verify_pptx(project, out)
    assert result["status"] == "PASS"


def test_real_pptx_inspection_with_libreoffice(tmp_path):
    import shutil
    import pytest
    from speaker_notes_builder.core import inspect_deck
    if not (shutil.which("soffice") or shutil.which("libreoffice")):
        pytest.skip("LibreOffice not available")
    pptx = tmp_path / "source.pptx"
    _make_pptx(pptx)
    project = tmp_path / "inspect-project"
    manifest = inspect_deck(pptx, project, profile="default", dpi=96)
    assert manifest["source_type"] == "pptx"
    assert manifest["slide_count"] == 2
    assert len(list((project / "slides").glob("slide-*.png"))) == 2
