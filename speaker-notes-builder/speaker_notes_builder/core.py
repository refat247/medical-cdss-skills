from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import subprocess
from difflib import SequenceMatcher
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import fitz
from PIL import Image
from pptx import Presentation


def sha256_file(path: str | Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def dhash_image(path: str | Path, hash_size: int = 8) -> str:
    img = Image.open(path).convert("L").resize((hash_size + 1, hash_size))
    px = list(img.get_flattened_data()) if hasattr(img, "get_flattened_data") else list(img.getdata())
    bits = []
    for y in range(hash_size):
        row = px[y * (hash_size + 1):(y + 1) * (hash_size + 1)]
        bits.extend(row[x] > row[x + 1] for x in range(hash_size))
    value = 0
    for bit in bits:
        value = (value << 1) | int(bit)
    width = hash_size * hash_size // 4
    return f"{value:0{width}x}"


def hamming_hex(a: str, b: str) -> int:
    return (int(a, 16) ^ int(b, 16)).bit_count()


def guess_role(title: str, text: str, word_count: int) -> str:
    title_l = title.lower().strip()
    t = (title + " " + text).lower()
    if any(k in title_l for k in ["reference", "bibliography"]):
        return "reference"
    if any(k in title_l for k in ["case 1", "case 2", "clinical case", "case study"]):
        return "case"
    if any(k in title_l for k in ["mechanism"]):
        return "mechanism" if word_count >= 14 else "divider"
    if any(k in title_l for k in ["dosing", "titration", "dose limit"]):
        return "dosing" if word_count >= 14 else "divider"
    if any(k in title_l for k in ["safety"]):
        return "safety" if word_count >= 14 else "divider"
    if any(k in title_l for k in ["clinical application", "regulatory status", "efficacy", "why it matters"]):
        return "divider"
    if any(k in title_l for k in ["roadmap", "learning objectives"]):
        return "general"
    if any(k in t for k in ["contraindication", "boxed warning", "monitoring", "hypogly"]):
        return "safety"
    if any(k in t for k in ["missed dose", "maintenance dose", "titration", "maximum dose"]):
        return "dosing"
    if any(k in t for k in ["counselling", "counseling", "contraception", "pregnan"]):
        return "counselling"
    if any(k in t for k in ["mechanism", "receptor", "agonism", "glp-1", "gip"]):
        return "mechanism"
    if any(k in t for k in ["surpass", "surmount", "hba1c", "trial", "placebo", "semaglutide"]):
        return "evidence"
    if any(k in t for k in ["takeaway", "remember", "checklist", "summary", "conclusion"]):
        return "takeaway"
    if word_count < 14:
        return "divider"
    return "general"


def semantic_density(word_count: int) -> str:
    if word_count < 25:
        return "low"
    if word_count < 75:
        return "medium"
    return "high"


def extract_title(text: str) -> str:
    lines = [re.sub(r"\s+", " ", x).strip() for x in text.splitlines() if x.strip()]
    for line in lines[:6]:
        if len(line) <= 120 and not re.fullmatch(r"\d+", line):
            return line
    return lines[0] if lines else "Untitled slide"


def _copy_source(input_path: Path, project: Path) -> Path:
    source_dir = project / "source"
    source_dir.mkdir(parents=True, exist_ok=True)
    dst = source_dir / f"source{input_path.suffix.lower()}"
    if input_path.resolve() != dst.resolve():
        shutil.copy2(input_path, dst)
    return dst


def _render_pdf(pdf_path: Path, slides_dir: Path, dpi: int = 150) -> list[Path]:
    slides_dir.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(pdf_path)
    out = []
    scale = dpi / 72.0
    matrix = fitz.Matrix(scale, scale)
    for i, page in enumerate(doc):
        pix = page.get_pixmap(matrix=matrix, alpha=False)
        path = slides_dir / f"slide-{i+1:03d}.png"
        pix.save(str(path))
        out.append(path)
    doc.close()
    return out


def _pptx_to_pdf(pptx_path: Path, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("LibreOffice/soffice is required to render PPTX slides.")
    cmd = [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(pptx_path)]
    proc = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=120)
    candidate = out_dir / (pptx_path.stem + ".pdf")
    if not candidate.exists():
        raise RuntimeError(f"PPTX render failed. stdout={proc.stdout!r} stderr={proc.stderr!r}")
    return candidate


def _extract_pdf_text(pdf_path: Path) -> list[str]:
    doc = fitz.open(pdf_path)
    texts = [page.get_text("text") or "" for page in doc]
    doc.close()
    return texts


def _extract_pptx_text_and_notes(pptx_path: Path) -> tuple[list[str], list[str]]:
    prs = Presentation(str(pptx_path))
    texts: list[str] = []
    notes: list[str] = []
    for slide in prs.slides:
        parts = []
        for shape in slide.shapes:
            try:
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    parts.append(shape.text.strip())
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        for cell in row.cells:
                            if cell.text.strip():
                                parts.append(cell.text.strip())
            except Exception:
                pass
        texts.append("\n".join(parts))
        note_text = ""
        try:
            tf = slide.notes_slide.notes_text_frame
            note_text = tf.text.strip() if tf else ""
        except Exception:
            note_text = ""
        notes.append(note_text)
    return texts, notes


def inspect_deck(input_path: str | Path, project_dir: str | Path, profile: str = "default", dpi: int = 150) -> dict[str, Any]:
    input_path = Path(input_path).resolve()
    project = Path(project_dir).resolve()
    project.mkdir(parents=True, exist_ok=True)
    if input_path.suffix.lower() not in {".pptx", ".pdf"}:
        raise ValueError("Input must be .pptx or .pdf")

    copied = _copy_source(input_path, project)
    source_hash = sha256_file(copied)
    slides_dir = project / "slides"
    render_dir = project / "render"

    if copied.suffix.lower() == ".pdf":
        render_pdf = copied
        texts = _extract_pdf_text(copied)
        notes = [""] * len(texts)
        slide_images = _render_pdf(copied, slides_dir, dpi=dpi)
        source_type = "pdf"
    else:
        texts, notes = _extract_pptx_text_and_notes(copied)
        render_pdf = _pptx_to_pdf(copied, render_dir)
        slide_images = _render_pdf(render_pdf, slides_dir, dpi=dpi)
        source_type = "pptx"
        if len(slide_images) != len(texts):
            raise RuntimeError(f"Rendered slide count {len(slide_images)} != PPTX slide count {len(texts)}")

    slide_records = []
    hashes: list[str] = []
    for i, (img, text, note) in enumerate(zip(slide_images, texts, notes), start=1):
        words = re.findall(r"\b\w+[\w\-/.%]*\b", text)
        title = extract_title(text)
        ph = dhash_image(img)
        hashes.append(ph)
        slide_records.append({
            "slide": i,
            "title": title,
            "text": text,
            "existing_notes": note,
            "image": str(img.relative_to(project)),
            "word_count": len(words),
            "role_guess": "title" if i == 1 else guess_role(title, text, len(words)),
            "semantic_density": semantic_density(len(words)),
            "image_sha256": sha256_file(img),
            "perceptual_hash": ph,
            "is_duplicate_of": None,
            "include": True,
        })

    duplicates = []
    for i in range(len(slide_records)):
        for j in range(i):
            distance = hamming_hex(hashes[i], hashes[j])
            exact = slide_records[i]["image_sha256"] == slide_records[j]["image_sha256"]
            ti = re.sub(r"\s+", " ", slide_records[i].get("text", "")).strip().lower()
            tj = re.sub(r"\s+", " ", slide_records[j].get("text", "")).strip().lower()
            text_similarity = SequenceMatcher(None, ti, tj).ratio() if (ti or tj) else 1.0
            # Perceptual hashes alone can collide on sparse divider slides.
            # Require near-identical text as a second signal unless the raster is byte-identical.
            near_duplicate = distance <= 1 and text_similarity >= 0.96
            if exact or near_duplicate:
                slide_records[i]["is_duplicate_of"] = j + 1
                duplicates.append({
                    "slide": i + 1,
                    "duplicate_of": j + 1,
                    "exact": exact,
                    "perceptual_hamming": distance,
                    "text_similarity": round(text_similarity, 4),
                })
                break

    manifest = {
        "skill_version": "1.0.0",
        "profile": profile,
        "source_type": source_type,
        "source_file": str(copied.relative_to(project)),
        "source_sha256": source_hash,
        "source_original_name": input_path.name,
        "render_authority": "source PDF" if source_type == "pdf" else "LibreOffice-rendered PDF",
        "slide_count": len(slide_records),
        "slides": slide_records,
        "duplicates": duplicates,
    }
    (project / "deck_manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    preflight = {
        "status": "PASS_WITH_WARNINGS" if duplicates else "PASS",
        "source_sha256": source_hash,
        "slide_count": len(slide_records),
        "rendered_images": len(slide_images),
        "duplicates_detected": duplicates,
        "warnings": ["Duplicate pages detected; disposition must be explicit before final timing/build."] if duplicates else [],
    }
    (project / "preflight_report.json").write_text(json.dumps(preflight, indent=2), encoding="utf-8")
    return manifest


ROLE_WEIGHTS = {
    "divider": 0.35,
    "title": 0.45,
    "reference": 0.45,
    "takeaway": 0.8,
    "mechanism": 1.0,
    "evidence": 1.25,
    "safety": 1.30,
    "dosing": 1.30,
    "counselling": 1.25,
    "case": 1.45,
    "general": 1.0,
}
DENSITY_MULTIPLIER = {"low": 0.8, "medium": 1.0, "high": 1.25}


def allocate_timing(project_dir: str | Path, talk_minutes: float, qa_minutes: float = 0.0) -> dict[str, Any]:
    project = Path(project_dir)
    manifest = json.loads((project / "deck_manifest.json").read_text(encoding="utf-8"))
    available = max(0.0, (talk_minutes - qa_minutes) * 60.0)
    included = [s for s in manifest["slides"] if s.get("include", True)]
    if not included:
        raise ValueError("No included slides")
    weights = []
    for s in included:
        w = ROLE_WEIGHTS.get(s.get("role_guess", "general"), 1.0)
        w *= DENSITY_MULTIPLIER.get(s.get("semantic_density", "medium"), 1.0)
        if s.get("is_duplicate_of"):
            w *= 0.35
        weights.append(w)
    total_w = sum(weights)
    raw = [available * w / total_w for w in weights]
    rounded = [max(4, int(round(x))) for x in raw]
    delta = int(round(available - sum(rounded)))
    idx = 0
    step = 1 if delta > 0 else -1
    while delta != 0 and rounded:
        k = idx % len(rounded)
        if step > 0 or rounded[k] > 4:
            rounded[k] += step
            delta -= step
        idx += 1
        if idx > 100000:
            break
    slides = []
    for s, seconds in zip(included, rounded):
        slides.append({
            "slide": s["slide"],
            "role": s.get("role_guess"),
            "density": s.get("semantic_density"),
            "allocated_seconds": seconds,
        })
    plan = {
        "requested_talk_minutes": talk_minutes,
        "qa_minutes": qa_minutes,
        "available_speaking_seconds": int(round(available)),
        "allocated_seconds": sum(x["allocated_seconds"] for x in slides),
        "slides": slides,
    }
    (project / "timing_plan.json").write_text(json.dumps(plan, indent=2), encoding="utf-8")
    return plan
