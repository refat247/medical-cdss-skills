from __future__ import annotations

import json
import shutil
import subprocess
from pathlib import Path
from typing import Any

import fitz
from pptx import Presentation


def validate_project(project_dir: str | Path) -> dict[str, Any]:
    project = Path(project_dir)
    problems = []
    warnings = []
    manifest = json.loads((project / "deck_manifest.json").read_text(encoding="utf-8")) if (project / "deck_manifest.json").exists() else None
    scripts = json.loads((project / "scripts.json").read_text(encoding="utf-8")) if (project / "scripts.json").exists() else None
    briefs = json.loads((project / "slide_briefs.json").read_text(encoding="utf-8")) if (project / "slide_briefs.json").exists() else None
    timing = json.loads((project / "timing_plan.json").read_text(encoding="utf-8")) if (project / "timing_plan.json").exists() else None

    if not manifest:
        problems.append("deck_manifest.json missing")
    if not scripts:
        problems.append("scripts.json missing")
    if manifest and scripts:
        included_manifest = {int(s["slide"]) for s in manifest["slides"] if s.get("include", True)}
        script_slides = [int(s["slide"]) for s in scripts.get("slides", []) if s.get("included", True)]
        if len(script_slides) != len(set(script_slides)):
            problems.append("duplicate script slide numbers")
        missing = sorted(included_manifest - set(script_slides))
        extra = sorted(set(script_slides) - included_manifest)
        if missing:
            problems.append(f"missing scripts for slides {missing}")
        if extra:
            warnings.append(f"scripts exist for non-included slides {extra}")
        for rec in scripts.get("slides", []):
            if rec.get("included", True) and not rec.get("main_script", "").strip():
                problems.append(f"slide {rec.get('slide')} has empty main_script")
    if briefs and manifest:
        brief_slides = {int(x["slide"]) for x in briefs.get("slides", [])}
        expected = {int(s["slide"]) for s in manifest["slides"] if s.get("include", True)}
        missing = sorted(expected - brief_slides)
        if missing:
            problems.append(f"missing slide briefs for slides {missing}")
        unresolved = []
        for b in briefs.get("slides", []):
            for c in b.get("coverage", []):
                if c.get("disposition") == "unresolved":
                    unresolved.append((b.get("slide"), c.get("element")))
        if unresolved:
            problems.append(f"unresolved semantic coverage items: {unresolved[:10]}")
    else:
        warnings.append("slide_briefs.json missing; semantic coverage not certified")
    if not (project / "deck_analysis.json").exists():
        warnings.append("deck_analysis.json missing; deck-first analysis not certified")
    if manifest and manifest.get("profile") == "clinical_cme":
        ledger_path = project / "claim_ledger.json"
        if not ledger_path.exists():
            problems.append("clinical_cme profile requires claim_ledger.json")
        else:
            ledger = json.loads(ledger_path.read_text(encoding="utf-8"))
            bad = [c for c in ledger.get("claims", []) if c.get("status") in {"UNVERIFIED", "BLOCKED"}]
            if bad:
                problems.append(f"claim ledger contains {len(bad)} UNVERIFIED/BLOCKED high-risk claim(s)")
    if timing and scripts:
        est = sum(float(x.get("estimated_seconds", 0)) for x in scripts.get("slides", []) if x.get("included", True))
        target = float(timing.get("available_speaking_seconds", 0))
        if target and abs(est - target) > 30:
            warnings.append(f"script estimated time {est:.0f}s differs from target {target:.0f}s by >30s")

    report = {"status": "FAIL" if problems else ("PASS_WITH_WARNINGS" if warnings else "PASS"), "problems": problems, "warnings": warnings}
    (project / "validation_report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    return report


def verify_docx(docx_path: str | Path, output_dir: str | Path, project_dir: str | Path | None = None) -> dict[str, Any]:
    docx_path = Path(docx_path).resolve()
    out = Path(output_dir).resolve()
    out.mkdir(parents=True, exist_ok=True)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        result = {"status": "UNCERTIFIED", "reason": "LibreOffice not available"}
        (out / "verify_docx.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
        return result
    proc = subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out), str(docx_path)], stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=120)
    pdf = out / f"{docx_path.stem}.pdf"
    if not pdf.exists():
        result = {"status": "FAIL", "reason": "DOCX to PDF conversion failed", "stdout": proc.stdout, "stderr": proc.stderr}
        (out / "verify_docx.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
        return result
    doc = fitz.open(pdf)
    pages = []
    for i, page in enumerate(doc):
        pix = page.get_pixmap(matrix=fitz.Matrix(1.2, 1.2), alpha=False)
        p = out / f"page-{i+1:03d}.png"
        pix.save(str(p))
        pages.append(str(p))
    count = len(doc)
    doc.close()
    status = "PASS"
    problems = []
    expected_pages = None
    if project_dir is not None:
        project = Path(project_dir)
        scripts_path = project / "scripts.json"
        if scripts_path.exists():
            scripts = json.loads(scripts_path.read_text(encoding="utf-8"))
            expected_pages = sum(1 for x in scripts.get("slides", []) if x.get("included", True))
            if count != expected_pages:
                status = "FAIL"
                problems.append(f"rendered page count {count} != expected presenter pages {expected_pages}; possible overflow or pagination defect")
    result = {"status": status, "page_count": count, "expected_pages": expected_pages, "problems": problems, "pdf": str(pdf), "page_images": pages, "note": "Render verification checks pagination when a project is supplied; visual inspection of every page remains required."}
    (out / "verify_docx.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    return result


def inject_notes(project_dir: str | Path, output_path: str | Path) -> Path:
    project = Path(project_dir)
    manifest = json.loads((project / "deck_manifest.json").read_text(encoding="utf-8"))
    if manifest.get("source_type") != "pptx":
        raise ValueError("Notes injection requires a PPTX source")
    scripts = json.loads((project / "scripts.json").read_text(encoding="utf-8"))
    source = project / manifest["source_file"]
    prs = Presentation(str(source))
    mapping = {int(x["slide"]): x for x in scripts.get("slides", []) if x.get("included", True)}
    for i, slide in enumerate(prs.slides, start=1):
        if i not in mapping:
            continue
        text = mapping[i].get("main_script", "").strip()
        slide.notes_slide.notes_text_frame.text = text
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.resolve() == source.resolve():
        raise ValueError("Refusing to overwrite source PPTX")
    prs.save(str(out))
    return out


def verify_pptx(project_dir: str | Path, pptx_path: str | Path) -> dict[str, Any]:
    project = Path(project_dir)
    manifest = json.loads((project / "deck_manifest.json").read_text(encoding="utf-8"))
    scripts = json.loads((project / "scripts.json").read_text(encoding="utf-8"))
    expected = {int(x["slide"]): x.get("main_script", "").strip() for x in scripts.get("slides", []) if x.get("included", True)}
    prs = Presentation(str(pptx_path))
    problems = []
    if len(prs.slides) != int(manifest["slide_count"]):
        problems.append(f"slide count changed: {len(prs.slides)} != {manifest['slide_count']}")
    checked = 0
    for i, slide in enumerate(prs.slides, start=1):
        if i not in expected:
            continue
        try:
            got = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            got = ""
        if got != expected[i]:
            problems.append(f"slide {i} notes mismatch")
        checked += 1
    result = {"status": "FAIL" if problems else "PASS", "slides_checked": checked, "problems": problems}
    qa_dir = project / "qa"; qa_dir.mkdir(exist_ok=True)
    (qa_dir / "verify_pptx.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    return result


def build_quality_report(project_dir: str | Path) -> dict[str, Any]:
    project = Path(project_dir)
    validation = json.loads((project / "validation_report.json").read_text(encoding="utf-8")) if (project / "validation_report.json").exists() else {"status": "UNCERTIFIED", "warnings": ["validation_report.json missing"]}
    clinical = json.loads((project / "clinical_lint.json").read_text(encoding="utf-8")) if (project / "clinical_lint.json").exists() else {"status": "UNCERTIFIED", "warnings": ["clinical_lint.json missing"]}
    preflight = json.loads((project / "preflight_report.json").read_text(encoding="utf-8")) if (project / "preflight_report.json").exists() else {"status": "UNCERTIFIED"}
    resolved_warnings = []
    effective_preflight = preflight.get("status")
    manifest_path = project / "deck_manifest.json"
    if manifest_path.exists() and preflight.get("duplicates_detected"):
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        by_slide = {int(x["slide"]): x for x in manifest.get("slides", [])}
        unresolved_dup = []
        for d in preflight.get("duplicates_detected", []):
            rec = by_slide.get(int(d.get("slide", 0)), {})
            if rec.get("duplicate_disposition"):
                resolved_warnings.append(f"Slide {d.get('slide')} duplicate disposition: {rec.get('duplicate_disposition')}")
            else:
                unresolved_dup.append(d)
        if not unresolved_dup and preflight.get("status") == "PASS_WITH_WARNINGS":
            # Preserve raw preflight evidence but consider the duplicate warning closed.
            effective_preflight = "PASS"

    statuses = [effective_preflight, validation.get("status"), clinical.get("status")]
    if "FAIL" in statuses:
        status = "FAIL"
    elif "UNCERTIFIED" in statuses:
        status = "UNCERTIFIED"
    elif "PASS_WITH_WARNINGS" in statuses:
        status = "PASS_WITH_WARNINGS"
    else:
        status = "PASS"
    report = {
        "status": status,
        "preflight": preflight,
        "effective_preflight_status": effective_preflight,
        "resolved_warnings": resolved_warnings,
        "validation": validation,
        "clinical": clinical,
        "artifact_note": "DOCX visual inspection and PPTX verification are separate gates when those outputs are requested.",
    }
    (project / "quality_report.json").write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    md = [f"# Speaker Notes Builder Quality Report\n", f"**Status:** {status}\n", "## Preflight", f"- {preflight.get('status')}", "## Validation", f"- {validation.get('status')}", "## Clinical lint", f"- {clinical.get('status')}"]
    if validation.get("problems"):
        md += ["\n### Problems"] + [f"- {x}" for x in validation["problems"]]
    warnings = list(validation.get("warnings", [])) + list(clinical.get("warnings", [])) + list(preflight.get("warnings", []))
    if warnings:
        md += ["\n### Warnings"] + [f"- {x}" for x in warnings]
    (project / "quality_report.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    return report
