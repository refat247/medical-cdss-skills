from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

GREEN = "0E4B42"
ACCENT = "1E6D5B"
PALE = "EAF4F1"
ORANGE = "B85B32"
PALE_ORANGE = "F9EEE8"
DARK = "20322F"
MUTED = "5F6E6B"


def _shade(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def _set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def _no_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def _set_font(run, name="Aptos", size=10.5, bold=False, color=DARK, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is not None:
        rFonts.set(qn("w:eastAsia"), name)


def _add_label_paragraph(cell, label: str, value: str = "", accent: str = ACCENT, body_size: float = 9.3):
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label.upper())
    _set_font(r, size=7.5, bold=True, color=accent)
    if value:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(4)
        r2 = p2.add_run(value)
        _set_font(r2, size=body_size, color=DARK)
    return p


def _setup_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.45)
    sec.bottom_margin = Inches(0.45)
    sec.left_margin = Inches(0.5)
    sec.right_margin = Inches(0.5)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    return doc


def _header_table(doc: Document, slide: int, section: str, seconds: float, mode: str):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.45, 5.25, 0.95]
    for c, w in zip(table.rows[0].cells, widths):
        c.width = Inches(w)
        _shade(c, GREEN)
        _set_cell_margins(c, 70, 100, 70, 100)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    left, mid, right = table.rows[0].cells
    p = left.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(f"SLIDE {slide}"); _set_font(r, size=8.5, bold=True, color="FFFFFF")
    p = mid.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run((section or "PRESENTATION").upper()); _set_font(r, size=8.5, bold=True, color="FFFFFF")
    p = right.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(f"~{int(round(seconds))}s"); _set_font(r, size=8.5, bold=True, color="FFFFFF")
    _no_table_borders(table)
    return table


def _add_slide_image(cell, image_path: Path, width_in: float):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(str(image_path), width=Inches(width_in))


def build_docx(project_dir: str | Path, mode: str = "rehearsal", output_path: str | Path | None = None) -> Path:
    project = Path(project_dir)
    manifest = json.loads((project / "deck_manifest.json").read_text(encoding="utf-8"))
    scripts = json.loads((project / "scripts.json").read_text(encoding="utf-8"))
    script_by_slide = {int(x["slide"]): x for x in scripts.get("slides", []) if x.get("included", True)}

    final = project / "final"
    final.mkdir(exist_ok=True)
    if output_path is None:
        output = final / ("rehearsal.docx" if mode == "rehearsal" else "live-presenter.docx")
    else:
        output = Path(output_path)

    doc = _setup_document()
    first = True
    for s in manifest["slides"]:
        num = int(s["slide"])
        if num not in script_by_slide:
            continue
        rec = script_by_slide[num]
        if not first:
            doc.add_page_break()
        first = False
        _header_table(doc, num, rec.get("section") or s.get("role_guess", ""), rec.get("estimated_seconds", 0), mode)
        doc.add_paragraph().paragraph_format.space_after = Pt(1)

        layout = doc.add_table(rows=1, cols=2)
        layout.autofit = False
        layout.alignment = WD_TABLE_ALIGNMENT.CENTER
        left, right = layout.rows[0].cells
        if mode == "rehearsal":
            left.width = Inches(3.7); right.width = Inches(3.95)
            _add_slide_image(left, project / s["image"], 3.45)
            _add_label_paragraph(right, "Slide title", rec.get("title") or s.get("title", ""), body_size=12.5)
            _add_label_paragraph(right, "Delivery intent", rec.get("delivery_intent", ""), body_size=9.4)
            if rec.get("pause_or_point"):
                _add_label_paragraph(right, "Pause / point", " • ".join(rec.get("pause_or_point", [])), body_size=8.8)
        else:
            left.width = Inches(3.25); right.width = Inches(4.4)
            _add_slide_image(left, project / s["image"], 3.0)
            p = right.paragraphs[0]
            rr = p.add_run(rec.get("title") or s.get("title", "")); _set_font(rr, size=13.5, bold=True, color=GREEN)
            p2 = right.add_paragraph()
            r2 = p2.add_run(rec.get("delivery_intent", "")); _set_font(r2, size=8.6, italic=True, color=MUTED)
        _no_table_borders(layout)

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run("SPEAKER SCRIPT" if mode == "rehearsal" else "MAIN SCRIPT")
        _set_font(r, size=8, bold=True, color=ACCENT)
        body = doc.add_paragraph()
        body.paragraph_format.space_after = Pt(5)
        body.paragraph_format.line_spacing = 1.06
        rb = body.add_run(rec.get("main_script", ""))
        _set_font(rb, size=10.2 if mode == "rehearsal" else 11.1, color=DARK)

        emphasis = rec.get("key_emphasis") or []
        if emphasis:
            t = doc.add_table(rows=1, cols=1)
            c = t.cell(0, 0); _shade(c, PALE); _set_cell_margins(c, 70, 110, 70, 110)
            p = c.paragraphs[0]; r = p.add_run("KEY EMPHASIS  "); _set_font(r, size=7.6, bold=True, color=ACCENT)
            r = p.add_run(" • ".join(emphasis)); _set_font(r, size=9.0 if mode == "rehearsal" else 9.5, bold=True, color=GREEN)
            _no_table_borders(t)

        cautions = rec.get("cautions") or []
        if cautions:
            t = doc.add_table(rows=1, cols=1)
            c = t.cell(0, 0); _shade(c, PALE_ORANGE); _set_cell_margins(c, 70, 110, 70, 110)
            p = c.paragraphs[0]; r = p.add_run("CLINICAL / REGULATORY CAUTION  "); _set_font(r, size=7.6, bold=True, color=ORANGE)
            r = p.add_run(" • ".join(cautions)); _set_font(r, size=8.8, color=ORANGE)
            _no_table_borders(t)

        short = rec.get("compressed_script", "")
        if short:
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(1)
            r = p.add_run("SHORT ON TIME"); _set_font(r, size=7.6, bold=True, color=ACCENT)
            p2 = doc.add_paragraph(); p2.paragraph_format.space_after = Pt(3)
            r2 = p2.add_run(short); _set_font(r2, size=8.8 if mode == "rehearsal" else 9.3, italic=True, color=DARK)

        if mode == "rehearsal" and rec.get("optional_expansion"):
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(1)
            r = p.add_run("OPTIONAL EXPANSION"); _set_font(r, size=7.6, bold=True, color=MUTED)
            p2 = doc.add_paragraph(); p2.paragraph_format.space_after = Pt(3)
            r2 = p2.add_run(rec.get("optional_expansion", "")); _set_font(r2, size=8.6, color=MUTED)

        trans = rec.get("transition", "")
        if trans:
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(0)
            r = p.add_run("NEXT  "); _set_font(r, size=7.5, bold=True, color=ACCENT)
            r2 = p.add_run(trans); _set_font(r2, size=8.8, italic=True, color=GREEN)

    doc.save(output)
    return output
