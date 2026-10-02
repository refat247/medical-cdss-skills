#!/usr/bin/env python3
"""
pptx_preflight.py -- renderer-free structural, typographic, geometric and
accessibility preflight for .pptx decks.

WHY THIS EXISTS
---------------
Render-based slide QA (PPTX -> PDF -> raster -> look at it) is the standard
approach, but it has two hard failure modes:

  1. It needs a converter (LibreOffice/soffice). Many sandboxes do not have one.
  2. Even when present, the converter substitutes missing fonts. Text metrics
     then differ from the user's PowerPoint, so a deck can pass the render
     gate and still overflow on the user's machine (and vice versa).

This tool therefore checks what can be checked *deterministically from the
OOXML itself*, and -- critically -- reports UNDETERMINED instead of PASS when a
check genuinely cannot be resolved without rendering.

USAGE
-----
    python3 pptx_preflight.py deck.pptx [--json out.json] [--min-pt 18]
            [--margin-in 0.4] [--expect-aspect 16:9] [--max-words 40]
            [--overflow-tol 1.02] [--fail-on error]

EXIT CODES
    0 = no ERROR-severity findings
    1 = at least one ERROR-severity finding
    2 = tool could not analyse the file at all

Requires: python-pptx (pip install python-pptx). Pillow optional but strongly
recommended -- without it, text-overflow estimation is skipped and reported as
UNDETERMINED rather than silently passed.
"""

import argparse
import json
import os
import re
import subprocess
import sys
import zipfile
import xml.etree.ElementTree as ET
from collections import defaultdict

EMU_PER_INCH = 914400          # ECMA-376 / ISO-IEC 29500 DrawingML unit
EMU_PER_PT = 12700

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}

# Fonts whose open substitutes are *metric-compatible* (same advance widths),
# so wrapping/overflow measured with the substitute transfers to the real font.
METRIC_COMPATIBLE = {
    "arial": "Liberation Sans",
    "helvetica": "Liberation Sans",
    "times new roman": "Liberation Serif",
    "courier new": "Liberation Mono",
    "calibri": "Carlito",
    "cambria": "Caladea",
}

PLACEHOLDER_PAT = re.compile(
    r"(lorem\s+ipsum|\bxxxx+\b|\bTBD\b|\bTODO\b|click\s+to\s+add|"
    r"add\s+(?:your\s+)?(?:title|text|subtitle)\s+here|"
    r"this\s+(?:page|slide)\s+.{0,20}layout|sample\s+text|placeholder)",
    re.I,
)

FOOTNOTE_PAT = re.compile(
    r"(doi:|pmid|et al\.?|nejm|jama|lancet|bmj|cochrane|diabetes care|"
    r"n engl j med|references?|adapted from|source:|accessed|\b(19|20)\d{2}\b)",
    re.I,
)

SEV_ORDER = {"ERROR": 3, "WARN": 2, "ADVISORY": 1, "UNDETERMINED": 1, "INFO": 0}


# --------------------------------------------------------------------------- #
# finding collection
# --------------------------------------------------------------------------- #
class Report:
    def __init__(self):
        self.findings = []

    def add(self, gate, severity, message, slide=None, shape=None, data=None):
        self.findings.append(
            {
                "gate": gate,
                "severity": severity,
                "slide": slide,
                "shape": shape,
                "message": message,
                "data": data or {},
            }
        )

    def counts(self):
        c = defaultdict(int)
        for f in self.findings:
            c[f["severity"]] += 1
        return dict(c)

    def by_gate(self):
        g = defaultdict(list)
        for f in self.findings:
            g[f["gate"]].append(f)
        return g


# --------------------------------------------------------------------------- #
# font environment
# --------------------------------------------------------------------------- #
def installed_font_families():
    """Return lowercase set of locally installed font families via fontconfig."""
    try:
        out = subprocess.run(
            ["fc-list", ":", "family"], capture_output=True, text=True, timeout=20
        ).stdout
    except Exception:
        return None
    fams = set()
    for line in out.splitlines():
        for part in line.split(","):
            part = part.strip()
            if part:
                fams.add(part.lower())
    return fams


def font_file_for(family):
    try:
        out = subprocess.run(
            ["fc-match", "-f", "%{file}", family],
            capture_output=True, text=True, timeout=20,
        ).stdout.strip()
        return out or None
    except Exception:
        return None


def resolve_internal_target(rels_part, target):
    """Resolve an OPC relationship target to a package member path.

    Relationship targets starting with "/" are package-root relative. The first
    draft treated them as filesystem-root paths and therefore flagged valid
    notes/master/slide relationships as dangling.
    """
    if target.startswith("/"):
        return os.path.normpath(target.lstrip("/")).replace("\\", "/")
    base = os.path.dirname(os.path.dirname(rels_part))
    return os.path.normpath(os.path.join(base, target)).replace("\\", "/")


# --------------------------------------------------------------------------- #
# S: package / structural gates
# --------------------------------------------------------------------------- #
def gate_structure(path, rep):
    info = {}
    try:
        zf = zipfile.ZipFile(path)
    except Exception as e:
        rep.add("S1-zip", "ERROR", "File is not a readable ZIP/OPC package: %s" % e)
        return info, None

    bad = zf.testzip()
    if bad:
        rep.add("S1-zip", "ERROR", "Corrupt ZIP member (CRC failure): %s" % bad)
    else:
        rep.add("S1-zip", "INFO", "ZIP integrity OK (%d parts)" % len(zf.namelist()))

    names = set(zf.namelist())
    info["part_count"] = len(names)

    for required in ("[Content_Types].xml", "_rels/.rels", "ppt/presentation.xml"):
        if required not in names:
            rep.add("S2-opc", "ERROR", "Required OPC part missing: %s" % required)

    # --- slide list vs slide parts -----------------------------------------
    slide_parts = sorted(n for n in names
                         if re.fullmatch(r"ppt/slides/slide\d+\.xml", n))
    info["slide_parts"] = len(slide_parts)

    listed_rids, listed_count = [], 0
    try:
        pres = ET.fromstring(zf.read("ppt/presentation.xml"))
        lst = pres.find("p:sldIdLst", NS)
        if lst is not None:
            for sid in lst.findall("p:sldId", NS):
                listed_rids.append(sid.get("{%s}id" % NS["r"]))
            listed_count = len(listed_rids)

        sz = pres.find("p:sldSz", NS)
        if sz is not None:
            cx, cy = int(sz.get("cx")), int(sz.get("cy"))
            info["slide_cx_emu"], info["slide_cy_emu"] = cx, cy
            info["slide_w_in"] = round(cx / EMU_PER_INCH, 4)
            info["slide_h_in"] = round(cy / EMU_PER_INCH, 4)
            info["aspect"] = round(cx / cy, 4)
    except Exception as e:
        rep.add("S3-presentation", "ERROR",
                "ppt/presentation.xml unparseable: %s" % e)

    info["slides_in_sldIdLst"] = listed_count
    if listed_count and listed_count != len(slide_parts):
        rep.add("S4-sldidlst", "WARN",
                "sldIdLst references %d slides but package holds %d slide parts; "
                "unreferenced slide parts are dead weight and will not display"
                % (listed_count, len(slide_parts)),
                data={"listed": listed_count, "parts": len(slide_parts)})
    else:
        rep.add("S4-sldidlst", "INFO",
                "sldIdLst consistent with slide parts (%d)" % len(slide_parts))

    # --- relationship targets resolve --------------------------------------
    dangling = []
    for n in names:
        if not n.endswith(".rels"):
            continue
        base = os.path.dirname(os.path.dirname(n))
        try:
            root = ET.fromstring(zf.read(n))
        except Exception:
            rep.add("S6-rels", "ERROR", "Unparseable relationship part: %s" % n)
            continue
        for r in root.findall("rel:Relationship", NS):
            if (r.get("TargetMode") or "") == "External":
                continue
            tgt = r.get("Target") or ""
            resolved = resolve_internal_target(n, tgt)
            if resolved not in names:
                dangling.append((n, tgt))
    if dangling:
        rep.add("S6-rels", "ERROR",
                "%d relationship target(s) do not resolve to a package part "
                "(this is the classic 'PowerPoint found a problem with content' "
                "repair prompt)" % len(dangling),
                data={"examples": dangling[:6]})
    else:
        rep.add("S6-rels", "INFO", "All internal relationship targets resolve")

    # --- orphan media -------------------------------------------------------
    media = {n for n in names if n.startswith("ppt/media/")}
    referenced = set()
    for n in names:
        if n.endswith(".rels"):
            try:
                root = ET.fromstring(zf.read(n))
            except Exception:
                continue
            for r in root.findall("rel:Relationship", NS):
                tgt = r.get("Target") or ""
                if "media/" in tgt:
                    referenced.add(resolve_internal_target(n, tgt))
    orphan = media - referenced
    if orphan:
        rep.add("S7-media", "ADVISORY",
                "%d media part(s) present but unreferenced (bloats file size)"
                % len(orphan), data={"examples": sorted(orphan)[:5]})
    info["media_parts"] = len(media)
    return info, zf


# --------------------------------------------------------------------------- #
# per-slide analysis via python-pptx
# --------------------------------------------------------------------------- #
def analyse(path, rep, args, info):
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        rep.add("T0-deps", "UNDETERMINED",
                "python-pptx not installed; typographic/geometric/accessibility "
                "gates could not run. Install with: pip install python-pptx")
        return

    try:
        prs = Presentation(path)
    except Exception as e:
        rep.add("S8-open", "ERROR", "python-pptx cannot open the deck: %s" % e)
        return

    slide_w, slide_h = prs.slide_width, prs.slide_height
    installed = installed_font_families()

    try:
        from PIL import ImageFont
        have_pil = True
    except ImportError:
        have_pil = False
        rep.add("G4-overflow", "UNDETERMINED",
                "Pillow not installed; text-overflow estimation skipped. "
                "This is the single most valuable geometric gate -- install Pillow.")

    fonts_used = defaultdict(int)
    min_pt_seen = None
    titles = []
    autofit_stale = 0
    unresolved_size_runs = 0
    total_runs = 0

    for idx, slide in enumerate(prs.slides, start=1):
        boxes = []          # (name, l, t, r, b, has_text)
        prior_fills = []    # (l, t, r, b, rgb), in z-order before current shape
        words_on_slide = 0
        try:
            slide_bg = solid_rgb(slide.background.fill)
        except Exception:
            slide_bg = None

        # ---- slide title (navigation + assistive tech) --------------------
        title_txt = None
        try:
            if slide.shapes.title is not None:
                title_txt = (slide.shapes.title.text or "").strip()
        except Exception:
            pass
        if not title_txt:
            rep.add("C4-title", "WARN",
                    "Slide has no non-empty title placeholder; screen readers and "
                    "the slide outline lose their primary navigation label",
                    slide=idx)
        else:
            titles.append((idx, title_txt))

        for shp in slide.shapes:
            nm = getattr(shp, "name", "?")

            # ---------- geometry --------------------------------------------
            try:
                l, t = int(shp.left), int(shp.top)
                w, h = int(shp.width), int(shp.height)
            except Exception:
                rep.add("G0-geom", "UNDETERMINED",
                        "Shape '%s' has no resolvable position/size (likely "
                        "inherits from layout); geometry gates skipped for it"
                        % nm, slide=idx, shape=nm)
                continue
            r_, b_ = l + w, t + h
            shp_fill_rgb = solid_rgb(getattr(shp, "fill", None))

            if l < 0 or t < 0 or r_ > slide_w or b_ > slide_h:
                over = max(0, -l), max(0, -t), max(0, r_ - slide_w), max(0, b_ - slide_h)
                rep.add("G1-offslide", "ERROR",
                        "Shape '%s' extends beyond the slide canvas "
                        "(L/T/R/B overflow in inches: %.2f/%.2f/%.2f/%.2f)"
                        % ((nm,) + tuple(x / EMU_PER_INCH for x in over)),
                        slide=idx, shape=nm)
            else:
                m_in = args.margin_in * EMU_PER_INCH
                gaps = (l, t, slide_w - r_, slide_h - b_)
                if min(gaps) < m_in:
                    rep.add("G2-margin", "WARN",
                            "Shape '%s' sits %.2f in from the nearest slide edge "
                            "(threshold %.2f in)"
                            % (nm, min(gaps) / EMU_PER_INCH, args.margin_in),
                            slide=idx, shape=nm)

            has_text = bool(getattr(shp, "has_text_frame", False)
                            and (shp.text_frame.text or "").strip())
            boxes.append((nm, l, t, r_, b_, has_text))

            # ---------- alt text -------------------------------------------
            st = str(getattr(shp, "shape_type", "") or "")
            is_visual = ("PICTURE" in st or "CHART" in st or "TABLE" in st
                         or "DIAGRAM" in st or "MEDIA" in st)
            if is_visual:
                descr = ""
                try:
                    descr = (shp._element._nvXxPr.cNvPr.get("descr") or "")
                except Exception:
                    try:
                        descr = shp._element.nvSpPr.cNvPr.get("descr") or ""
                    except Exception:
                        descr = ""
                d = descr.strip()
                if not d:
                    rep.add("A1-alttext", "ERROR",
                            "Visual element '%s' (%s) has no alt text; fails "
                            "WCAG 1.1.1 non-text content" % (nm, st),
                            slide=idx, shape=nm)
                elif re.fullmatch(r"[\w \-]+\.(png|jpe?g|gif|bmp|tiff?|svg|emf|wmf)", d, re.I) \
                        or re.fullmatch(r"(image|picture|chart|graph|figure|table)\s*\d*", d, re.I):
                    # python-pptx and PowerPoint both auto-populate descr with the
                    # source filename. It satisfies a naive presence check while
                    # conveying nothing to a screen reader -- worse than empty,
                    # because it silently passes.
                    rep.add("A1-alttext", "ERROR",
                            "Visual element '%s' has auto-generated alt text %r "
                            "(filename or generic label). This passes a presence "
                            "check but conveys no information; WCAG 1.1.1 requires "
                            "a text alternative that serves the equivalent purpose"
                            % (nm, d), slide=idx, shape=nm)

            # ---------- table header row -----------------------------------
            if getattr(shp, "has_table", False):
                try:
                    if not shp.table.first_row:
                        rep.add("A2-tablehdr", "WARN",
                                "Table '%s' has no header row flagged; assistive "
                                "tech cannot associate cells with headers "
                                "(WCAG 1.3.1)" % nm, slide=idx, shape=nm)
                except Exception:
                    pass

            if not getattr(shp, "has_text_frame", False):
                if shp_fill_rgb:
                    prior_fills.append((l, t, r_, b_, shp_fill_rgb))
                continue

            tf = shp.text_frame
            txt = tf.text or ""
            words_on_slide += len(txt.split())
            text_role = classify_text_role(shp, txt, slide_h)

            # ---------- placeholder text -----------------------------------
            m = PLACEHOLDER_PAT.search(txt)
            if m:
                rep.add("C1-placeholder", "ERROR",
                        "Unreplaced placeholder/boilerplate text in '%s': %r"
                        % (nm, m.group(0)), slide=idx, shape=nm)

            # ---------- autofit --------------------------------------------
            bp = tf._txBody.find("a:bodyPr", NS)
            fontscale = None
            if bp is not None:
                na = bp.find("a:normAutofit", NS)
                if na is not None:
                    fs = na.get("fontScale")
                    if fs:
                        fontscale = int(fs) / 100000.0
                        autofit_stale += 1
                        rep.add("T3-autofit", "WARN",
                                "'%s' uses normAutofit with a stored fontScale of "
                                "%.1f%%. That value is a cache written by the "
                                "authoring app, not a recomputed value -- text "
                                "edited programmatically will NOT rescale, so the "
                                "rendered size here may not match PowerPoint's"
                                % (nm, fontscale * 100),
                                slide=idx, shape=nm,
                                data={"fontScale": fontscale})

            # ---------- runs: size + font ----------------------------------
            for para in tf.paragraphs:
                for run in para.runs:
                    total_runs += 1
                    sz = run.font.size
                    if sz is None:
                        unresolved_size_runs += 1
                    else:
                        pt = sz.pt
                        eff = pt * (fontscale or 1.0)
                        if min_pt_seen is None or eff < min_pt_seen:
                            min_pt_seen = eff
                        floor = args.citation_min_pt if text_role == "citation" else args.min_pt
                        if eff < floor:
                            rep.add("T1-minsize", "WARN",
                                    "Run in '%s' renders at ~%.1f pt (below the "
                                    "%.0f pt %s floor)" % (nm, eff, floor, text_role),
                                    slide=idx, shape=nm,
                                    data={"declared_pt": pt, "effective_pt": eff,
                                          "role": text_role})
                    nmf = run.font.name
                    if nmf:
                        fonts_used[nmf] += 1

                    # ---------- A3: WCAG 1.4.3 contrast --------------------
                    try:
                        fg = None
                        if _is_rgb(run.font.color):
                            fg = tuple(int(x) for x in
                                       bytes.fromhex(str(run.font.color.rgb)))
                    except Exception:
                        fg = None
                    if fg is None:
                        continue
                    bg = shp_fill_rgb or containing_prior_fill(prior_fills, l, t, r_, b_) or slide_bg
                    if bg is None:
                        # WCAG 2.2 contrast-ratio Note 4: "It is a failure if no
                        # background color is specified when the text color is
                        # specified, because the user's default background color
                        # is unknown and cannot be evaluated."
                        rep.add("A3-contrast", "UNDETERMINED",
                                "'%s' sets an explicit text colour but no "
                                "resolvable background (shape fill and slide "
                                "background both inherit). Contrast cannot be "
                                "evaluated -- WCAG 2.2 treats an unspecified "
                                "background as a failure condition" % nm,
                                slide=idx, shape=nm)
                        continue
                    ratio = contrast_ratio(fg, bg)
                    pt_eff = (sz.pt * (fontscale or 1.0)) if sz else None
                    # WCAG 2.2 glossary "large scale (text)": >=18pt, or >=14pt bold
                    large = bool(pt_eff and (pt_eff >= 18 or
                                             (pt_eff >= 14 and run.font.bold)))
                    need = 3.0 if large else 4.5
                    if ratio < need:
                        rep.add("A3-contrast", "ERROR",
                                "'%s' text #%02X%02X%02X on #%02X%02X%02X has a "
                                "contrast ratio of %.2f:1; WCAG 2.2 SC 1.4.3 (AA) "
                                "requires %.1f:1 for this size (%s text)"
                                % ((nm,) + fg + bg + (ratio, need,
                                   "large-scale" if large else "normal")),
                                slide=idx, shape=nm,
                                data={"ratio": round(ratio, 2), "required": need})

            # ---------- overflow estimate ----------------------------------
            if have_pil and txt.strip():
                est = estimate_overflow(shp, tf, fontscale, args)
                if est is None:
                    rep.add("G4-overflow", "UNDETERMINED",
                            "'%s' text fit could not be measured (inherited font "
                            "size or unavailable font metrics)" % nm,
                            slide=idx, shape=nm)
                elif text_role == "citation" and est["ratio"] <= args.citation_overflow_tol:
                    pass
                elif est["mode"] == "spAutoFit":
                    # Box grows to fit text. The stored height is a cache the
                    # authoring app wrote; programmatic text edits do NOT update
                    # it. So the risk is not clipped text -- it is a box that
                    # silently grows past the canvas or over its neighbours.
                    if est["ratio"] > args.overflow_tol:
                        grown_b = t + int(est["needed_pt"] * EMU_PER_PT) + tIns_of(tf)
                        off = grown_b > slide_h
                        rep.add("G4-autogrow", "ERROR" if off else "WARN",
                                "'%s' is set to spAutoFit (grow shape to fit text) "
                                "and its text now needs ~%.2fx the stored height. "
                                "PowerPoint will expand the box to ~%.2f in on open"
                                "%s. The stored height in the file is stale."
                                % (nm, est["ratio"],
                                   est["needed_pt"] / 72.0,
                                   ", pushing it past the bottom of the slide"
                                   if off else ""),
                                slide=idx, shape=nm, data=est)
                elif est["ratio"] > args.overflow_tol:
                    sev = "ERROR" if est["ratio"] > 1.15 else "WARN"
                    rep.add("G4-overflow", sev,
                            "'%s' text needs ~%.2fx its box height "
                            "(%d wrapped lines in space for ~%d) -- text will be "
                            "clipped or spill outside the shape (autofit=%s)"
                            % (nm, est["ratio"], est["lines"], est["fit_lines"],
                               est["mode"]),
                            slide=idx, shape=nm, data=est)

            if shp_fill_rgb:
                prior_fills.append((l, t, r_, b_, shp_fill_rgb))

        # ---------- overlap of text-bearing shapes --------------------------
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, b = boxes[i], boxes[j]
                if not (a[5] or b[5]):
                    continue
                ox = min(a[3], b[3]) - max(a[1], b[1])
                oy = min(a[4], b[4]) - max(a[2], b[2])
                if ox > 0 and oy > 0:
                    area = ox * oy
                    amin = min((a[3] - a[1]) * (a[4] - a[2]),
                               (b[3] - b[1]) * (b[4] - b[2])) or 1
                    frac = area / amin
                    if frac > 0.03:
                        rep.add("G3-overlap", "WARN",
                                "'%s' and '%s' overlap over %.0f%% of the smaller "
                                "shape; at least one carries text"
                                % (a[0], b[0], frac * 100),
                                slide=idx,
                                data={"overlap_fraction": round(frac, 3)})

        if words_on_slide > args.max_words:
            rep.add("C5-density", "ADVISORY",
                    "Slide carries %d words (reporting threshold %d). Not a "
                    "defect by itself -- but text-heavy slides read aloud are the "
                    "specific configuration multimedia-learning research finds "
                    "least effective" % (words_on_slide, args.max_words),
                    slide=idx, data={"words": words_on_slide})

    # ---------- duplicate titles -------------------------------------------
    seen = defaultdict(list)
    for i, t in titles:
        seen[t.lower()].append(i)
    for t, idxs in seen.items():
        if len(idxs) > 1:
            rep.add("C3-duptitle", "ADVISORY",
                    "Title %r repeats on slides %s; consider '(1/n)' continuation "
                    "labels so the outline stays navigable" % (t, idxs))

    # ---------- font substitution risk -------------------------------------
    info["fonts_used"] = dict(fonts_used)
    info["min_effective_pt"] = min_pt_seen
    info["runs_without_explicit_size"] = unresolved_size_runs
    info["runs_total"] = total_runs

    if unresolved_size_runs:
        rep.add("T4-inherit", "UNDETERMINED",
                "%d of %d text runs carry no explicit size and inherit from the "
                "layout/master. Their rendered size -- and therefore any overflow "
                "verdict -- cannot be resolved from the slide part alone"
                % (unresolved_size_runs, total_runs))

    if installed is None:
        rep.add("T2-fontsub", "UNDETERMINED",
                "fontconfig unavailable; cannot check which deck fonts exist here")
    else:
        for fam, n in sorted(fonts_used.items()):
            key = fam.lower()
            if key in installed:
                continue
            mc = METRIC_COMPATIBLE.get(key)
            if mc and mc.lower() in installed:
                rep.add("T2-fontsub", "ADVISORY",
                        "Font %r is not installed, but metric-compatible "
                        "substitute %r is -- wrapping measurements transfer"
                        % (fam, mc), data={"runs": n})
            else:
                actual = font_file_for(fam)
                rep.add("T2-fontsub", "ERROR",
                        "Font %r (%d runs) is NOT installed and has no "
                        "metric-compatible substitute here; fontconfig falls back "
                        "to %s. Any render or overflow measurement taken in this "
                        "environment uses different advance widths than the "
                        "author's PowerPoint, so it CANNOT certify fit"
                        % (fam, n, os.path.basename(actual) if actual else "?"),
                        data={"runs": n, "fallback": actual})


def insets_of(tf):
    """Return (lIns, rIns, tIns, bIns) in EMU, honouring OOXML defaults."""
    bp = tf._txBody.find("a:bodyPr", NS)
    lIns = rIns = 91440      # OOXML defaults: 0.1 in L/R
    tIns = bIns = 45720      # 0.05 in T/B
    if bp is not None:
        for attr in ("lIns", "rIns", "tIns", "bIns"):
            v = bp.get(attr)
            if v is None:
                continue
            if attr == "lIns":   lIns = int(v)
            elif attr == "rIns": rIns = int(v)
            elif attr == "tIns": tIns = int(v)
            else:                bIns = int(v)
    return lIns, rIns, tIns, bIns


def tIns_of(tf):
    return insets_of(tf)[2]


def classify_text_role(shp, text, slide_h):
    """Best-effort role classification for thresholds, not content semantics."""
    try:
        top = int(shp.top)
        height = int(shp.height)
    except Exception:
        top = height = 0
    lower_band = top > (slide_h * 0.82)
    short = len((text or "").split()) <= 28
    if FOOTNOTE_PAT.search(text or "") or (lower_band and short):
        return "citation"
    return "body"


def containing_prior_fill(prior_fills, l, t, r_, b_):
    """Return the topmost earlier solid fill under the text's upper anchor.

    Safety-expanded text boxes can be taller than their visible first line; a
    center-point test then misses ordinary "label text over card/header"
    designs. The upper anchor is a better proxy for where the rendered text
    begins in these generated medical decks.
    """
    cx, cy = (l + r_) / 2.0, t + min((b_ - t) * 0.25, 160000)
    for fl, ft, fr, fb, rgb in reversed(prior_fills):
        if fl <= cx <= fr and ft <= cy <= fb:
            return rgb
    return None


def autofit_mode(tf):
    bp = tf._txBody.find("a:bodyPr", NS)
    if bp is None:
        return "unspecified"
    if bp.find("a:spAutoFit", NS) is not None:
        return "spAutoFit"
    if bp.find("a:normAutofit", NS) is not None:
        return "normAutofit"
    if bp.find("a:noAutofit", NS) is not None:
        return "noAutofit"
    return "unspecified"


def estimate_overflow(shp, tf, fontscale, args):
    """Wrap text with real font metrics and compare required vs available height.

    Returns dict with 'mode' so the caller can distinguish:
      noAutofit/unspecified -> text clips or spills (classic overflow)
      normAutofit           -> app shrinks text; stored fontScale is a stale cache
      spAutoFit             -> app grows the box; stored height is a stale cache
    """
    from PIL import ImageFont
    try:
        w = int(shp.width)
        h = int(shp.height)
    except Exception:
        return None

    lIns, rIns, tIns, bIns = insets_of(tf)
    mode = autofit_mode(tf)

    avail_w_pt = max(1.0, (w - lIns - rIns) / EMU_PER_PT)
    avail_h_pt = max(1.0, (h - tIns - bIns) / EMU_PER_PT)

    total_lines = 0
    max_pt = 0.0
    for para in tf.paragraphs:
        runs = [r for r in para.runs if (r.text or "")]
        if not runs:
            total_lines += 1
            continue
        pt = None
        fam = None
        for r in runs:
            if r.font.size is not None:
                pt = r.font.size.pt
            if r.font.name:
                fam = r.font.name
            if pt:
                break
        if pt is None:
            return None          # inherited size -> undeterminable, stay honest
        pt = pt * (fontscale or 1.0)
        max_pt = max(max_pt, pt)
        path = font_file_for(fam or "sans-serif")
        try:
            f = ImageFont.truetype(path, size=max(1, int(round(pt))))
        except Exception:
            return None
        text = "".join(r.text or "" for r in runs)
        # greedy wrap in points (PIL measures in px == pt at size=pt)
        words, line, lines = text.split(), "", 0
        for word in words:
            trial = (line + " " + word).strip()
            if f.getlength(trial) <= avail_w_pt or not line:
                line = trial
            else:
                lines += 1
                line = word
        lines += 1 if line else 0
        total_lines += max(1, lines)

    line_h = max_pt * 1.2                      # typical single line spacing
    need = total_lines * line_h
    fit = int(avail_h_pt // line_h) if line_h else 0
    return {
        # ECMA-376-1 S21.1.2.1.2: "If this element is omitted, then noAutofit
        # or auto-fit off is implied" -- so 'unspecified' behaves as noAutofit.
        "mode": "noAutofit" if mode == "unspecified" else mode,
        "lines": total_lines,
        "fit_lines": fit,
        "needed_pt": round(need, 1),
        "avail_pt": round(avail_h_pt, 1),
        "ratio": round(need / avail_h_pt, 3) if avail_h_pt else 99,
        "measured_with": os.path.basename(font_file_for(fam or "sans-serif") or "?"),
    }


# --------------------------------------------------------------------------- #
# A3: WCAG contrast
# --------------------------------------------------------------------------- #
def rel_luminance(rgb):
    """WCAG 2.2 relative luminance."""
    def chan(c):
        c = c / 255.0
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = (chan(x) for x in rgb)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def contrast_ratio(fg, bg):
    """WCAG 2.2: (L1 + 0.05) / (L2 + 0.05), lighter over darker."""
    l1, l2 = rel_luminance(fg), rel_luminance(bg)
    if l1 < l2:
        l1, l2 = l2, l1
    return (l1 + 0.05) / (l2 + 0.05)


def _is_rgb(color):
    """True only for an explicit sRGB value (not a theme colour, not None)."""
    try:
        return color is not None and color.type is not None \
            and getattr(color.type, "name", str(color.type)).startswith("RGB")
    except Exception:
        return False


def solid_rgb(fill):
    """Return (r,g,b) if the fill is an explicit solid RGB, else None.

    Theme colours and inherited/background fills return None on purpose -- they
    cannot be resolved without the theme part, and guessing would turn an
    UNDETERMINED contrast result into a false PASS.
    """
    try:
        if fill is None or int(fill.type) != 1:      # 1 == MSO_FILL_TYPE.SOLID
            return None
        c = fill.fore_color
        if not _is_rgb(c):
            return None
        return tuple(int(x) for x in bytes.fromhex(str(c.rgb)))
    except Exception:
        return None


# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx")
    ap.add_argument("--json")
    ap.add_argument("--min-pt", type=float, default=18.0)
    ap.add_argument("--citation-min-pt", type=float, default=10.0)
    ap.add_argument("--margin-in", type=float, default=0.4)
    ap.add_argument("--max-words", type=int, default=40)
    ap.add_argument("--overflow-tol", type=float, default=1.02)
    ap.add_argument("--citation-overflow-tol", type=float, default=1.18)
    ap.add_argument("--expect-aspect", default=None, help='e.g. "16:9"')
    ap.add_argument("--fail-on", default="error", choices=["error", "warn"])
    args = ap.parse_args()

    if not os.path.exists(args.pptx):
        print("FATAL: no such file: %s" % args.pptx, file=sys.stderr)
        return 2

    rep = Report()
    info, zf = gate_structure(args.pptx, rep)
    if zf is None:
        emit(rep, info, args)
        return 2

    if args.expect_aspect and info.get("aspect"):
        a, b = args.expect_aspect.split(":")
        want = float(a) / float(b)
        if abs(info["aspect"] - want) > 0.01:
            rep.add("S5-aspect", "ERROR",
                    "Slide aspect ratio is %.4f (%.2f x %.2f in) but %s was "
                    "required. Merging or presenting a mismatched deck letterboxes "
                    "or crops every slide"
                    % (info["aspect"], info.get("slide_w_in", 0),
                       info.get("slide_h_in", 0), args.expect_aspect))

    analyse(args.pptx, rep, args, info)
    emit(rep, info, args)

    worst = max((SEV_ORDER[f["severity"]] for f in rep.findings), default=0)
    if args.fail_on == "error":
        return 1 if worst >= 3 else 0
    return 1 if worst >= 2 else 0


def emit(rep, info, args):
    counts = rep.counts()
    print("=" * 78)
    print("PPTX PREFLIGHT  --  %s" % os.path.basename(args.pptx))
    print("=" * 78)
    print("slides in sldIdLst : %s" % info.get("slides_in_sldIdLst"))
    print("slide parts        : %s" % info.get("slide_parts"))
    if info.get("slide_w_in"):
        print("canvas             : %.2f x %.2f in (aspect %.4f)"
              % (info["slide_w_in"], info["slide_h_in"], info["aspect"]))
    print("media parts        : %s" % info.get("media_parts"))
    if info.get("min_effective_pt") is not None:
        print("smallest text      : %.1f pt (effective)" % info["min_effective_pt"])
    if info.get("fonts_used"):
        print("fonts referenced   : %s"
              % ", ".join("%s(%d)" % (k, v) for k, v in sorted(info["fonts_used"].items())))
    print("-" * 78)

    order = ["ERROR", "WARN", "UNDETERMINED", "ADVISORY", "INFO"]
    for sev in order:
        rows = [f for f in rep.findings if f["severity"] == sev]
        if not rows:
            continue
        print("\n### %s (%d)" % (sev, len(rows)))
        for f in rows:
            loc = ""
            if f["slide"]:
                loc = " [slide %s%s]" % (f["slide"],
                                         "/" + f["shape"] if f["shape"] else "")
            print("  - (%s)%s %s" % (f["gate"], loc, f["message"]))

    print("\n" + "-" * 78)
    print("SUMMARY: " + "  ".join("%s=%d" % (k, counts.get(k, 0)) for k in order))
    verdict = "FAIL" if counts.get("ERROR") else (
        "PASS-WITH-WARNINGS" if counts.get("WARN") else "PASS")
    if counts.get("UNDETERMINED"):
        verdict += " (+UNDETERMINED items -- not certified)"
    print("VERDICT: %s" % verdict)
    print("-" * 78)

    if args.json:
        with open(args.json, "w") as fh:
            json.dump({"file": args.pptx, "info": info,
                       "counts": counts, "verdict": verdict,
                       "findings": rep.findings}, fh, indent=1, default=str)
        print("JSON written: %s" % args.json)


if __name__ == "__main__":
    sys.exit(main())
