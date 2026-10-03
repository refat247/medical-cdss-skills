"""Generate fixtures/visual_polish/case_reveal_mini.pptx.

A synthetic 6-slide deck that mirrors the builder's case/reveal geometry
conventions (eyebrow, title, progress bar, 1-3 cards with case id + badge +
body, source line, speaker notes). Clinical-looking text is placeholder only
("FIXTURE ...") so the fixture can never be mistaken for evidence.

    python fixtures/visual_polish/make_fixture.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

PX = 9525  # 1280 x 720 px canvas, as the builder emits
OUT = Path(__file__).with_name('case_reveal_mini.pptx')


def box(slide, x, y, w, h, fill=None, text=None, size=16, bold=False):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
    if fill:
        sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sp.fill.background()
    sp.line.fill.background()
    if text is not None:
        tf = sp.text_frame
        tf.text = text
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(size); r.font.bold = bold
                r.font.color.rgb = RGBColor.from_string('17212B')
    return sp


def case_slide(prs, eyebrow, title, cases, ratio):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s, 53, 34, 595, 27, text=eyebrow, size=18, bold=True)
    box(s, 53, 69, 1138, 58, text=title, size=36, bold=True)
    box(s, 56, 140, 1166, 7, 'DDE5ED'); box(s, 56, 140, int(1166 * ratio), 7, '175CD3')
    k = len(cases); cw = {1: 1166, 2: 572, 3: 373}[k]; gap = (1166 - k * cw) / max(1, k - 1) if k > 1 else 0
    for i, (cid, badge, body) in enumerate(cases):
        x = int(56 + i * (cw + gap))
        box(s, x, 165, cw, 440, 'FFFFFF'); box(s, x, 165, 9, 440, '175CD3')
        box(s, x + 21, 182, 117, 36, 'EAF2FF'); box(s, x + 28, 185, 102, 27, text=cid, size=17, bold=True)
        box(s, x + cw - 119, 182, 96, 36, 'FFF5E6'); box(s, x + cw - 111, 185, 81, 27, text=badge, size=17, bold=True)
        box(s, x + 13, 236, cw - 26, 357, text=body, size=24)
    box(s, 56, 621, 1166, 24, text='SOURCE A: FIXTURE GUIDELINE', size=16.5)
    s.notes_slide.notes_text_frame.text = f'FIXTURE NOTES for {title}'
    return s


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(1280 * PX), Emu(720 * PX)
    case_slide(prs, 'CASE / DECIDE', 'DECIDE — Fixture Module',
               [('CASE 001', 'DECIDE', 'FIXTURE scenario one. Which decision applies?'),
                ('CASE 002', 'DECIDE', 'FIXTURE scenario two with x\u2082 subscript. Which decision applies?'),
                ('CASE 003', 'DECIDE', 'FIXTURE scenario three. Which decision applies?')], 0.1)
    case_slide(prs, 'GUIDELINE / SOURCE REVEAL', 'SOURCE REVEAL — Fixture Module',
               [('CASE 001', 'REVEAL • 1/B-NR', 'FIXTURE reveal one is recommended.'),
                ('CASE 002', 'REVEAL • 2a/B-R', 'FIXTURE reveal two is reasonable.'),
                ('CASE 003', 'REVEAL • 3: No Benefit/A; 1/C-LD', 'FIXTURE reveal three.')], 0.1)
    case_slide(prs, 'GUIDELINE / SOURCE REVEAL', 'SOURCE REVEAL — Fixture Single',
               [('CASE 004', 'REVEAL • 2b/C-LD', 'FIXTURE single reveal body text.')], 0.2)
    case_slide(prs, 'CASE / DECIDE', 'DECIDE — Fixture Overflow',
               [('CASE 005', 'DECIDE', ' '.join(['FIXTURE overflow sentence that is deliberately very long.'] * 14)),
                ('CASE 006', 'DECIDE', 'FIXTURE short.'),
                ('CASE 007', 'DECIDE', 'FIXTURE short.')], 0.3)
    s = prs.slides.add_slide(prs.slide_layouts[6])  # unknown role -> must be SKIPPED
    box(s, 53, 69, 1138, 58, text='FIXTURE free-form slide', size=36, bold=True)
    box(s, 53, 200, 1138, 300, text='FIXTURE body', size=24)
    # case slide with an extra unknown text box -> must be SKIPPED
    sl = case_slide(prs, 'CASE / DECIDE', 'DECIDE — Fixture Extra', [('CASE 008', 'DECIDE', 'FIXTURE body.')], 0.4)
    box(sl, 900, 30, 300, 30, text='FIXTURE stray label', size=16)
    prs.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
