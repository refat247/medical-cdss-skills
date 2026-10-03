"""Advisory geometry/text-fit screen for CASE/DECIDE and SOURCE REVEAL slides.

This detects likely collisions before rendering. It cannot certify PowerPoint
layout: fonts, text metrics, and Office line breaking vary by renderer.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation

def _metric_font():
    """Portable bold sans metric font (was hard-coded to C:/Windows/Fonts/arialbd.ttf)."""
    import os, shutil, subprocess
    cands = [os.environ.get('LAYOUT_GATE_FONT', ''),
             os.path.join(os.environ.get('WINDIR', 'C:/Windows'), 'Fonts', 'arialbd.ttf'),
             '/Library/Fonts/Arial Bold.ttf', '/System/Library/Fonts/Supplemental/Arial Bold.ttf',
             '/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf',
             '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf']
    for c in cands:
        if c and Path(c).exists():
            return Path(c)
    if shutil.which('fc-match'):
        out = subprocess.run(['fc-match', '-f', '%{file}', 'Arial:bold'], capture_output=True, text=True).stdout.strip()
        if out and Path(out).exists():
            return Path(out)
    raise FileNotFoundError('no bold sans-serif metric font found; set LAYOUT_GATE_FONT')


FONT = _metric_font()
EMU_PER_INCH = 914400
PX_PER_INCH = 96


def inch(value):
    return value / EMU_PER_INCH


def text_size(shape):
    sizes = [run.font.size.pt for p in shape.text_frame.paragraphs
             for run in p.runs if run.font.size]
    return max(sizes, default=24)


def estimated_lines(text, width_in, size_pt):
    """Word-wrap estimate using a common Windows sans-serif substitute."""
    font = ImageFont.truetype(str(FONT), max(1, round(size_pt * 96 / 72)))
    width_px = max(1, width_in * PX_PER_INCH)
    lines = 0
    for paragraph in text.split('\n'):
        words = paragraph.split()
        if not words:
            lines += 1
            continue
        current = ''
        for word in words:
            candidate = f'{current} {word}'.strip()
            if current and font.getlength(candidate) > width_px:
                lines += 1
                current = word
            else:
                current = candidate
        lines += 1
    return lines


def issue(slide, code, detail):
    return {'slide': slide, 'code': code, 'detail': detail}


def audit(path, slide_numbers=None):
    deck = Presentation(path)
    selected = set(slide_numbers) if slide_numbers else set(range(1, len(deck.slides) + 1))
    issues = []
    audited = []
    for number in sorted(selected):
        if number < 1 or number > len(deck.slides):
            issues.append(issue(number, 'OUT_OF_RANGE', 'slide is absent from deck'))
            continue
        slide = deck.slides[number - 1]
        text_shapes = [s for s in slide.shapes if s.has_text_frame and s.text.strip()]
        titles = [s for s in text_shapes if inch(s.top) < 1.45 and
                  ('SOURCE REVEAL' in s.text.upper() or s.text.upper().startswith('DECIDE'))]
        if not titles:
            continue
        audited.append(number)
        title = min(titles, key=lambda s: inch(s.top))
        # Slide 5-like decks have a separate small "GUIDELINE / SOURCE REVEAL"
        # eyebrow. Choose the larger title, not the eyebrow.
        title = max(titles, key=lambda s: text_size(s))
        title_lines = estimated_lines(title.text, inch(title.width) - 0.12, text_size(title))
        title_height = title_lines * text_size(title) / 72 * 1.12
        if title_height > inch(title.height) + 0.08:
            issues.append(issue(number, 'TITLE_TEXT_FIT',
                                f'estimated {title_lines} lines at {text_size(title):g} pt exceed title box'))
        eyebrows = [s for s in text_shapes if s is not title and inch(s.top) < 0.85
                    and ('GUIDELINE' in s.text.upper() or 'CASE / DECIDE' in s.text.upper())]
        if eyebrows and inch(title.top) < max(inch(s.top + s.height) for s in eyebrows) - 0.03:
            issues.append(issue(number, 'TITLE_EYEBROW_COLLISION', 'title box intersects eyebrow'))
        badges = [s for s in text_shapes if s.text.strip().upper().startswith('REVEAL')]
        bodies = [s for s in text_shapes if 2.25 <= inch(s.top) <= 3.0 and inch(s.height) >= 2.5]
        footer = [s for s in text_shapes if inch(s.top) >= 6.3 and
                  ('SOURCE' in s.text.upper() or 'GUIDELINE' in s.text.upper())]
        footer_top = min((inch(s.top) for s in footer), default=7.5)
        for badge in badges:
            lines = estimated_lines(badge.text, inch(badge.width) - 0.08, text_size(badge))
            height = lines * text_size(badge) / 72 * 1.08
            if height > inch(badge.height) + 0.04:
                issues.append(issue(number, 'BADGE_TEXT_FIT', f'{badge.text[:36]!r} likely clips or wraps'))
            if any(inch(badge.top + badge.height) > inch(body.top) + 0.03 and
                   inch(badge.left) < inch(body.left + body.width) and
                   inch(body.left) < inch(badge.left + badge.width) for body in bodies):
                issues.append(issue(number, 'BADGE_BODY_COLLISION', 'reveal badge intersects body box'))
        for body in bodies:
            if inch(body.top + body.height) > footer_top - 0.12:
                issues.append(issue(number, 'BODY_FOOTER_COLLISION', 'body box enters source/footer region'))
    return {'gate': 'CASE_REVEAL_LAYOUT_ADVISORY', 'pptx': str(path),
            'slide_count': len(deck.slides), 'audited_slide_numbers': audited,
            'issues': issues, 'pass': not issues,
            'certifies_visual_quality': False}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('pptx')
    parser.add_argument('--slides', nargs='*', type=int)
    args = parser.parse_args()
    result = audit(args.pptx, args.slides)
    print(json.dumps(result, indent=2))
    raise SystemExit(0 if result['pass'] else 1)


if __name__ == '__main__':
    main()
