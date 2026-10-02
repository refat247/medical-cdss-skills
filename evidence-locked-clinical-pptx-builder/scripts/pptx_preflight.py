from __future__ import annotations
import argparse, json, math, re, zipfile
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

PLACEHOLDER_PATTERNS=[r'\bTBD\b',r'\bTODO\b',r'\bPLACEHOLDER\b',r'undefined',r'\{\{.+?\}\}']
def text_runs(shape):
    if not getattr(shape,'has_text_frame',False): return []
    out=[]
    for p in shape.text_frame.paragraphs:
        for r in p.runs: out.append(r)
    return out
def rect(shape): return (shape.left,shape.top,shape.left+shape.width,shape.top+shape.height)
def inter(a,b):
    x=max(0,min(a[2],b[2])-max(a[0],b[0])); y=max(0,min(a[3],b[3])-max(a[1],b[1])); return x*y
def area(r): return max(0,r[2]-r[0])*max(0,r[3]-r[1])
def audit(path,min_body=24.0,min_source=16.0,expected_ratio=16/9):
    issues=[]; warnings=[]
    try:
        with zipfile.ZipFile(path) as zf:
            bad=zf.testzip()
            if bad: issues.append({'issue':'OOXML zip corruption','part':bad})
            names=set(zf.namelist())
            if 'ppt/presentation.xml' not in names: issues.append({'issue':'missing ppt/presentation.xml'})
    except zipfile.BadZipFile:
        return {'file':str(path),'pass':False,'issues':[{'issue':'not a valid PPTX ZIP'}],'warnings':[]}
    prs=Presentation(path); sw,slide_h=prs.slide_width,prs.slide_height; ratio=sw/slide_h
    if abs(ratio-expected_ratio)>0.05: warnings.append({'issue':'unexpected aspect ratio','ratio':ratio})
    min_font=None; small_body=[]; small_source=[]; bounds=[]; overlap=[]; placeholders=[]; low_dpi=[]; crops=[]; title_missing=[]
    for si,slide in enumerate(prs.slides,1):
        text_shapes=[]; has_title=False
        for sh in slide.shapes:
            if sh.left<0 or sh.top<0 or sh.left+sh.width>sw or sh.top+sh.height>slide_h:
                bounds.append({'slide':si,'shape':getattr(sh,'name','')})
            if getattr(sh,'has_text_frame',False):
                txt=(sh.text or '').strip()
                # title heuristic: non-empty text in top quarter
                if txt and sh.top < slide_h*0.25: has_title=True
                text_shapes.append(sh)
                for run in text_runs(sh):
                    if run.font.size is None: continue
                    pt=run.font.size.pt; min_font=pt if min_font is None else min(min_font,pt)
                    t=(run.text or '').strip()
                    source_like=bool(re.search(r'ESC|ACC|AHA|guideline|source|doi|UDMI|reference',t,re.I))
                    if source_like and pt < min_source: small_source.append({'slide':si,'text':t[:60],'pt':pt})
                    elif t and pt < min_body: small_body.append({'slide':si,'text':t[:60],'pt':pt})
                for pat in PLACEHOLDER_PATTERNS:
                    if re.search(pat,txt,re.I): placeholders.append({'slide':si,'text':txt[:100],'pattern':pat})
            if sh.shape_type==MSO_SHAPE_TYPE.PICTURE:
                try:
                    img=sh.image; dpi_x=img.size[0]/(sh.width/914400); dpi_y=img.size[1]/(sh.height/914400); dpi=min(dpi_x,dpi_y)
                    if dpi<96: low_dpi.append({'slide':si,'shape':sh.name,'effective_dpi':round(dpi,1)})
                    crop=sum(getattr(sh,x,0) or 0 for x in ['crop_left','crop_right','crop_top','crop_bottom'])
                    if crop>0.7: crops.append({'slide':si,'shape':sh.name,'crop_sum':round(crop,3)})
                except Exception: pass
        if not has_title and si>1: title_missing.append(si)
        for i,a in enumerate(text_shapes):
            ra=rect(a)
            if area(ra)==0: continue
            for b in text_shapes[i+1:]:
                rb=rect(b); ia=inter(ra,rb)
                if ia and ia/min(area(ra),area(rb))>0.20: overlap.append({'slide':si,'a':a.name,'b':b.name,'overlap_ratio':round(ia/min(area(ra),area(rb)),3)})
    if bounds: issues.append({'issue':'out_of_bounds_shapes','items':bounds})
    if small_source: issues.append({'issue':'source_label_font_below_floor','items':small_source})
    if small_body: warnings.append({'issue':'body_runs_below_default_floor','items':small_body})
    if overlap: warnings.append({'issue':'text_box_overlap_heuristic','items':overlap})
    if placeholders: issues.append({'issue':'placeholder_or_undefined_text','items':placeholders})
    if low_dpi: warnings.append({'issue':'low_effective_image_dpi','items':low_dpi})
    if crops: warnings.append({'issue':'heavy_image_crop','items':crops})
    if title_missing: warnings.append({'issue':'title_missing_heuristic','slides':title_missing})
    return {'file':str(path),'slides':len(prs.slides),'aspect_ratio':ratio,'min_detected_font_pt':min_font,'issues':issues,'warnings':warnings,'pass':not issues,'note':'Automated preflight cannot certify visual fit/projector quality; final render review remains mandatory.'}
def main():
    ap=argparse.ArgumentParser(); ap.add_argument('pptx'); ap.add_argument('--min-body',type=float,default=24); ap.add_argument('--min-source',type=float,default=16); a=ap.parse_args(); r=audit(a.pptx,a.min_body,a.min_source); print(json.dumps(r,indent=2)); raise SystemExit(1 if not r['pass'] else 0)
if __name__=='__main__': main()
