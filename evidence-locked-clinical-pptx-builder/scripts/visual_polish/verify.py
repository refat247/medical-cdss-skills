#!/usr/bin/env python3
"""VISUAL_POLISH_CONTENT_LOCK gate: proves a polished deck says exactly what the built deck said.

    python -m scripts.visual_polish.verify BUILT.pptx POLISHED.pptx \
        [--polish-report polish_report.json] [--render OUTDIR] [--json visual_polish_verify.json]

Blocking checks (exit 1 on any failure):
  1. identity    - same slide count, same slide ids in the same order, same layout per slide
  2. text lock   - every original text atom is present verbatim on the same slide
                   (whitespace-normalised); multiset-checked so nothing is dropped
  3. no addition - every polished text box is original text or a whitelisted design token
  4. notes       - speaker-notes text byte-identical per slide
  5. chrome      - slide layouts / masters byte-identical (footer, brand, chrome untouched)
  6. floors      - no polished run below the source-label floor (16 pt); body runs not
                   below the profile floor recorded in the polish report
  7. fit         - each rebuilt body box re-measured independently, must fit
Then it runs the repo's own gates on the polished deck:
  scripts/pptx_preflight.py and scripts/case_reveal_layout_gate.py (advisory)
Optional --render: LibreOffice -> PNG per slide + render manifest (via
scripts/render_qa_gate.py) + a blank per-slide visual_review_ledger.csv with
the required domain columns, ready for the mandatory 100% individual review.

This gate never certifies visual quality and never replaces SKILL.md sections 18-20.
"""
from __future__ import annotations

import argparse
import csv
import glob
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from collections import Counter

from pptx import Presentation

try:
    from scripts.visual_polish import style as S
    from scripts import pptx_preflight, case_reveal_layout_gate
    from scripts.visual_review_ledger import DOMAIN_COLUMNS
except ModuleNotFoundError:  # direct execution
    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
    from scripts.visual_polish import style as S
    from scripts import pptx_preflight, case_reveal_layout_gate
    from scripts.visual_review_ledger import DOMAIN_COLUMNS

# Exactly the non-source text the polish stage may introduce.
TOKEN_RE = [r'\?', '\u2713', r'REVEAL\s*•', '\u2192', r'\d{1,3}']
SOURCE_FLOOR = 16.0
BADGE_SHAPES = {'Case id', 'Phase badge', 'COR/LOE badge', 'Reveal label'}


def norm(s):
    return re.sub(r'\s+', ' ', (s or '').replace('\v', ' ')).strip()


def slide_texts(slide):
    return [sp.text_frame.text for sp in slide.shapes
            if getattr(sp, 'has_text_frame', False) and sp.text_frame.text.strip()]


def notes_text(slide):
    if not slide.has_notes_slide or slide.notes_slide.notes_text_frame is None:
        return ''
    return slide.notes_slide.notes_text_frame.text


def atoms(t):
    """Split an original element into the units the polish stage may lay out separately."""
    t = t.strip()
    if t.startswith('REVEAL'):
        code = re.sub(r'^REVEAL\s*•?\s*', '', t)
        return [norm(p) for p in code.split(';') if p.strip()] or ['REVEAL']
    if '\u2192' in t:
        return [norm(p) for p in re.split(r'[\u2192\n]', t) if p.strip()]
    if '\n' in t or '\v' in t:
        return [norm(p) for p in re.split(r'[\n\v]', t) if p.strip()]
    return [norm(t)]


def is_token(s):
    s = norm(s)
    for tok in TOKEN_RE:
        s = re.sub(rf'(?<!\w){tok}(?!\w)', ' ', s)
    return not norm(s)


def chrome_parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()
                if re.match(r'ppt/(slideLayouts|slideMasters|theme)/[^/]+\.xml$', n)}


def check_content(orig_path, pol_path):
    o, p = Presentation(orig_path), Presentation(pol_path)
    errs = []
    if len(o.slides) != len(p.slides):
        return [f'slide count changed {len(o.slides)} -> {len(p.slides)}']
    for i, (a, b) in enumerate(zip(o.slides, p.slides), 1):
        if a.slide_id != b.slide_id:
            errs.append(f'slide {i}: slide id/order changed ({a.slide_id} -> {b.slide_id})')
        if a.slide_layout.name != b.slide_layout.name:
            errs.append(f'slide {i}: layout changed')
        ta, tb = slide_texts(a), slide_texts(b)
        want = Counter(x for t in ta for x in atoms(t))
        have_boxes = [norm(t) for t in tb]
        have_join = ' \u241e '.join(have_boxes)
        for atom, n in want.items():
            if have_join.count(atom) < n:
                errs.append(f'slide {i}: MISSING original text (x{n - have_join.count(atom)}): {atom[:90]!r}')
        orig_join = ' \u241e '.join(norm(t) for t in ta)
        for t in tb:
            nt = norm(t)
            if nt in orig_join or is_token(nt):
                continue
            # a polished box may hold a sub-part of an original element (split badge/flow/title)
            if any(nt in want_atom or want_atom in nt for want_atom in want) and all(w in orig_join for w in nt.split()):
                continue
            errs.append(f'slide {i}: ADDED text not in built slide: {nt[:90]!r}')
        if notes_text(a) != notes_text(b):
            errs.append(f'slide {i}: speaker notes changed')
    oc, pc = chrome_parts(orig_path), chrome_parts(pol_path)
    if oc != pc:
        changed = sorted(k for k in set(oc) | set(pc) if oc.get(k) != pc.get(k))
        errs.append(f'layout/master/theme chrome changed: {changed[:5]}')
    return errs


def check_floors_and_fit(pol_path, polish_report):
    p = Presentation(pol_path)
    S.set_canvas(p)
    k = S.DESIGN_W / (p.slide_width / 12700)
    by_slide = {r['slide']: r for r in (polish_report or {}).get('slides_report', [])}
    prof = S.PROFILES.get((polish_report or {}).get('profile', 'projector_default'), S.PROFILES['projector_default'])
    errs = []
    for i, sl in enumerate(p.slides, 1):
        r = by_slide.get(i)
        if r is not None and r.get('status') != 'POLISHED':
            continue  # left as built; judged by the repo's normal preflight
        for sp in sl.shapes:
            if not getattr(sp, 'has_text_frame', False) or not sp.text_frame.text.strip():
                continue
            # real pt -> design pt (k == 1.0 on 13.333 in decks)
            sizes = [run.font.size.pt * k for para in sp.text_frame.paragraphs
                     for run in para.runs if run.font.size]
            if not sizes:
                continue
            mn = min(sizes)
            body = sp.name.startswith('Body ')
            badge = sp.name in BADGE_SHAPES
            floor_lbl = S.BADGE_MIN if badge else SOURCE_FLOOR
            if not body and mn < floor_lbl - 0.05 and not is_token(sp.text_frame.text):
                errs.append(f'slide {i}: {sp.name!r} {mn:.1f} pt below {floor_lbl:g} pt floor')
            if body:
                cards = (r or {}).get('cards', 3)
                floor = prof['body_floor'].get(cards, 24)
                if mn < floor - 0.05:
                    errs.append(f'slide {i}: {sp.name!r} {mn:.1f} pt below {floor} pt body floor')
                w, h = sp.width / 12700 * k, sp.height / 12700 * k
                need = S.text_height(sp.text_frame.text, w, mn)
                if need > h + 1:
                    errs.append(f'slide {i}: {sp.name} needs {need:.0f} pt, box {h:.0f} pt')
    return errs


def render(pptx, outdir):
    os.makedirs(outdir, exist_ok=True)
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice:
        return dict(status='RENDER-UNCERTIFIED', reason='LibreOffice not found')
    ver = subprocess.run([soffice, '--version'], capture_output=True, text=True).stdout.strip()
    tmp = tempfile.mkdtemp()
    subprocess.run([soffice, '--headless', '--convert-to', 'pdf', '--outdir', tmp, pptx],
                   check=True, capture_output=True, timeout=1200)
    pdf = glob.glob(os.path.join(tmp, '*.pdf'))[0]
    try:
        import pymupdf as fitz
    except ImportError:
        try:
            import fitz
        except ImportError:
            return dict(status='RENDER-UNCERTIFIED', reason='pymupdf not installed', pdf=pdf, renderer=ver)
    doc = fitz.open(pdf)
    for n, page in enumerate(doc, 1):
        page.get_pixmap(dpi=144).save(os.path.join(outdir, f'slide_{n:03d}.png'))
    try:
        from scripts.render_qa_gate import gate
    except ModuleNotFoundError:
        from render_qa_gate import gate
    rg = gate(pptx, outdir, manifest_out=os.path.join(outdir, 'render_manifest.json'))
    ledger = os.path.join(outdir, 'visual_review_ledger.csv')
    if not os.path.exists(ledger):
        with open(ledger, 'w', newline='', encoding='utf-8') as f:
            w = csv.writer(f)
            w.writerow(['slide', 'render_version', 'reviewer', 'final_status', *DOMAIN_COLUMNS, 'issues', 'repair_action'])
            for n in range(1, len(doc) + 1):
                w.writerow([n, ver, '', '', *[''] * len(DOMAIN_COLUMNS), '', ''])
    return dict(status='RENDERED_REVIEW_PENDING', renderer=ver, slides=len(doc), dir=outdir,
                rendered_slide_set_sha256=rg.get('rendered_slide_set_sha256'),
                ledger=ledger, note='Fill the ledger slide by slide; render_qa_gate --require-ledger certifies.')


def verify(orig, pol, polish_report=None, render_dir=None):
    content = check_content(orig, pol)
    floors = check_floors_and_fit(pol, polish_report)
    pre = pptx_preflight.audit(pol)
    try:
        layout = case_reveal_layout_gate.audit(pol)
    except Exception as e:  # advisory only
        layout = dict(issues=[], error=str(e))
    rend = render(pol, render_dir) if render_dir else dict(status='NOT_REQUESTED')
    passed = not content and not floors and pre['pass']
    return dict(gate='VISUAL_POLISH_CONTENT_LOCK', original=str(orig), polished=str(pol),
                content_lock=dict(passed=not content, errors=content),
                floors_and_fit=dict(passed=not floors, errors=floors),
                repo_preflight=dict(passed=pre['pass'], issues=pre['issues'],
                                    warning_counts={w['issue']: len(w.get('items', w.get('slides', []))) for w in pre['warnings']}),
                case_reveal_layout_advisory=dict(issue_count=len(layout.get('issues', [])),
                                                 codes=dict(Counter(i['code'] for i in layout.get('issues', []))),
                                                 error=layout.get('error')),
                render=rend, pass_=passed, certifies_visual_quality=False)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('original'); ap.add_argument('polished')
    ap.add_argument('--polish-report')
    ap.add_argument('--render', help='output dir for PNG renders, render manifest and review ledger')
    ap.add_argument('--json', default=None)
    a = ap.parse_args(argv)
    rep = json.load(open(a.polish_report, encoding='utf-8')) if a.polish_report else None
    r = verify(a.original, a.polished, rep, a.render)
    r['pass'] = r.pop('pass_')
    if a.json:
        with open(a.json, 'w', encoding='utf-8') as f:
            json.dump(r, f, indent=2)
    print(f"CONTENT LOCK   : {'PASS' if r['content_lock']['passed'] else 'FAIL'} ({len(r['content_lock']['errors'])})")
    print(f"FLOORS & FIT   : {'PASS' if r['floors_and_fit']['passed'] else 'FAIL'} ({len(r['floors_and_fit']['errors'])})")
    print(f"REPO PREFLIGHT : {'PASS' if r['repo_preflight']['passed'] else 'FAIL'} warnings={r['repo_preflight']['warning_counts']}")
    print(f"LAYOUT SCREEN  : {r['case_reveal_layout_advisory']['issue_count']} advisory issues {r['case_reveal_layout_advisory']['codes']}")
    print(f"RENDER         : {r['render']['status']}")
    for e in (r['content_lock']['errors'] + r['floors_and_fit']['errors'])[:40]:
        print('  -', e)
    for iss in r['repo_preflight']['issues'][:5]:
        print('  - preflight:', iss.get('issue'), str(iss.get('items', ''))[:200])
    return 0 if r['pass'] else 1


if __name__ == '__main__':
    sys.exit(main())
