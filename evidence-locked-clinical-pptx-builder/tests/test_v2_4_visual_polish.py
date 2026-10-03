"""v2.4.0 VISUAL_POLISH regression tests.

The polish stage may change presentation only. These tests prove the fail-closed
contract: verbatim text, notes and chrome preserved; unknown structure skipped;
floors never violated; tampering detected; promotion blocked without the gate.
"""
import copy
import json
from pathlib import Path

import pytest
import yaml
from pptx import Presentation
from pptx.util import Pt

from scripts.visual_polish import style as S
from scripts.visual_polish.polish import polish, badge_parts, plan_header
from scripts.visual_polish.verify import check_content, check_floors_and_fit, verify
from scripts.semantic_promotion_gate import evaluate
from scripts.mode_router import route
from scripts.validate_project import validate_project
from scripts.state_manager import validate as validate_state, next_state

ROOT = Path(__file__).parents[1]
MINI = ROOT / 'fixtures/visual_polish/case_reveal_mini.pptx'
GATES = ROOT / 'fixtures/dyslipidemia_regression/gates_all_pass.json'
DEFECTS_CLOSED = ROOT / 'fixtures/dyslipidemia_regression/defects_closed.csv'


@pytest.fixture()
def polished(tmp_path):
    out = tmp_path / 'polished.pptx'
    rep = polish(MINI, out, report_path=tmp_path / 'r.json')
    return out, rep


def _status(rep):
    return {r['slide']: r['status'] for r in rep['slides_report']}


def test_roles_and_fail_closed_statuses(polished):
    _, rep = polished
    st = _status(rep)
    assert st[1] == st[2] == st[3] == 'POLISHED'
    assert st[4] == 'SPLIT_REQUIRED'          # cannot fit at 24 pt floor -> not shrunk
    assert st[5] == 'SKIPPED'                 # unknown role
    assert st[6] == 'SKIPPED'                 # unrecognised extra text element
    assert rep['certifies_anything'] is False


def test_content_lock_passes_on_polished_deck(polished):
    out, _ = polished
    assert check_content(MINI, out) == []


def test_skipped_and_split_slides_are_byte_identical_shape_trees(polished):
    out, _ = polished
    a, b = Presentation(MINI), Presentation(out)
    for n in (4, 5, 6):
        assert a.slides[n - 1].shapes._spTree.xml == b.slides[n - 1].shapes._spTree.xml


def test_notes_ids_order_and_layout_chrome_preserved(polished):
    out, _ = polished
    a, b = Presentation(MINI), Presentation(out)
    assert [s.slide_id for s in a.slides] == [s.slide_id for s in b.slides]
    assert [s.notes_slide.notes_text_frame.text for s in a.slides] == \
           [s.notes_slide.notes_text_frame.text for s in b.slides]


def test_subscript_and_compound_cor_badges_survive(polished):
    out, _ = polished
    b = Presentation(out)
    texts2 = [sp.text_frame.text for sp in b.slides[1].shapes if sp.has_text_frame]
    assert '3: No Benefit/A' in texts2 and '1/C-LD' in texts2
    texts1 = ' '.join(sp.text_frame.text for sp in b.slides[0].shapes if sp.has_text_frame)
    assert 'x\u2082 subscript' in texts1


def test_dropped_text_is_detected(polished, tmp_path):
    out, _ = polished
    prs = Presentation(out)
    body = next(sp for sp in prs.slides[0].shapes if sp.name.startswith('Body '))
    body.text_frame.paragraphs[0].runs[0].text = body.text_frame.paragraphs[0].runs[0].text[:-12]
    bad = tmp_path / 'bad.pptx'; prs.save(bad)
    assert any('MISSING' in e for e in check_content(MINI, bad))


def test_mutated_threshold_is_detected(tmp_path, polished):
    out, _ = polished
    prs = Presentation(out)
    body = next(sp for sp in prs.slides[1].shapes if sp.name.startswith('Body '))
    r = body.text_frame.paragraphs[0].runs[0]
    r.text = r.text.replace('recommended', 'not recommended')
    bad = tmp_path / 'bad.pptx'; prs.save(bad)
    errs = check_content(MINI, bad)
    assert any('MISSING' in e for e in errs) or any('ADDED' in e for e in errs)


def test_added_clinical_text_is_detected(polished, tmp_path):
    out, _ = polished
    prs = Presentation(out)
    prs.slides[0].shapes.add_textbox(0, 0, 100000, 100000).text_frame.text = 'FIXTURE invented dose 40 mg'
    bad = tmp_path / 'bad.pptx'; prs.save(bad)
    assert any('ADDED' in e for e in check_content(MINI, bad))


def test_changed_notes_are_detected(polished, tmp_path):
    out, _ = polished
    prs = Presentation(out)
    prs.slides[2].notes_slide.notes_text_frame.text = 'edited'
    bad = tmp_path / 'bad.pptx'; prs.save(bad)
    assert any('notes' in e for e in check_content(MINI, bad))


def test_reordered_slides_are_detected(polished, tmp_path):
    out, _ = polished
    prs = Presentation(out)
    lst = prs.slides._sldIdLst
    first = lst[0]; lst.remove(first); lst.append(first)
    bad = tmp_path / 'bad.pptx'; prs.save(bad)
    assert any('order' in e for e in check_content(MINI, bad))


def test_body_below_profile_floor_is_detected(polished, tmp_path):
    out, rep = polished
    prs = Presentation(out)
    body = next(sp for sp in prs.slides[0].shapes if sp.name.startswith('Body '))
    for p in body.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(18)
    bad = tmp_path / 'bad.pptx'; prs.save(bad)
    assert any('body floor' in e for e in check_floors_and_fit(bad, rep))


def test_dense_profile_requires_recorded_approval(tmp_path):
    with pytest.raises(SystemExit):
        polish(MINI, tmp_path / 'x.pptx', profile_name='dense_case_reveal')
    rep = polish(MINI, tmp_path / 'x.pptx', profile_name='dense_case_reveal', approved_by='deck owner')
    assert rep['profile_approved_by'] == 'deck owner'


def test_profiles_never_lower_source_floor_below_16():
    for prof in S.PROFILES.values():
        assert prof['source'] >= 16 and prof['label'] >= 16
        assert min(prof['body_floor'].values()) >= 18
    assert S.PROFILES['projector_default']['body_floor'] == {1: 24, 2: 24, 3: 24}


def test_fit_size_returns_none_instead_of_shrinking():
    long = ' '.join(['word'] * 400)
    assert S.fit_size([long], 200, 100, 24, 24) is None


def test_badge_parsing_and_header_wrap():
    assert badge_parts('REVEAL • 3: No Benefit/A; 1/C-LD') == ('REVEAL •', ['3: No Benefit/A', '1/C-LD'])
    one = plan_header(dict(id='CASE 001', badge='REVEAL • 1/B-NR'), 800)
    assert one['rows'] == 1 and one['size'] == 16
    tight = plan_header(dict(id='CASE 001', badge='REVEAL • 3: No Benefit/B-NR'), 230)
    assert tight['rows'] >= 2 and tight['size'] >= S.BADGE_MIN


def test_verify_end_to_end_without_render(polished):
    out, rep = polished
    r = verify(MINI, out, rep)
    assert r['content_lock']['passed'] and r['floors_and_fit']['passed']
    assert r['certifies_visual_quality'] is False


def test_promotion_blocks_when_polish_applied_but_gate_missing(tmp_path):
    g = json.loads(GATES.read_text())
    g['visual_polish_applied'] = True
    p = tmp_path / 'g.json'; p.write_text(json.dumps(g))
    r = evaluate(DEFECTS_CLOSED, p)
    assert not r['pass'] and 'visual_polish_content_lock' in r['missing_gates']
    g['visual_polish_content_lock'] = 'WAIVED'; p.write_text(json.dumps(g))
    assert 'visual_polish_content_lock' in evaluate(DEFECTS_CLOSED, p)['failed_gates']
    g['visual_polish_content_lock'] = 'PASS'; p.write_text(json.dumps(g))
    assert evaluate(DEFECTS_CLOSED, p)['pass']


def test_promotion_unchanged_when_polish_not_applied():
    assert evaluate(DEFECTS_CLOSED, GATES)['pass']


def test_state_template_has_visual_polish_between_audit_and_preflight():
    s = yaml.safe_load((ROOT / 'templates/project_state.yaml').read_text())
    assert validate_state(s) == []
    st = {x['id']: x for x in s['pipelines']['presentation_artifact']['states']}
    assert st['visual_polish']['depends_on'] == ['clinical_provenance_audit']
    assert st['automated_preflight']['depends_on'] == ['visual_polish']
    assert next_state(s)['id'] == 'evidence_lock'


def test_project_validator_requires_approval_for_dense_profile(tmp_path):
    s = yaml.safe_load((ROOT / 'templates/project_state.yaml').read_text())
    s['project']['id'] = 't'; s['project']['mode'] = 'VISUAL_POLISH'
    s['release_policy']['visual_polish'].update(enabled=True, profile='dense_case_reveal', profile_approved_by='')
    p = tmp_path / 's.yaml'; p.write_text(yaml.safe_dump(s, sort_keys=False))
    assert any('profile_approved_by' in e for e in validate_project(p))
    s['release_policy']['visual_polish']['profile_approved_by'] = 'owner'
    p.write_text(yaml.safe_dump(s, sort_keys=False))
    assert validate_project(p) == []


def test_mode_router_routes_beautify_requests():
    assert route('beautify this deck without changing content')['mode'] == 'VISUAL_POLISH'
    assert route('improve the visual of the pptx')['mode'] == 'VISUAL_POLISH'
    assert route('run visual qa on the deck')['mode'] == 'VISUAL_QA'
